import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import openpyxl, base64, os

# ── PAGE CONFIG ───────────────────────────────────────────────
st.set_page_config(
    page_title="Opponent Wise | Cricket Analysis",
    page_icon="", layout="wide",
    initial_sidebar_state="collapsed",
)

# ── PATHS ─────────────────────────────────────────────────────
BASE_DIR    = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR    = os.path.dirname(BASE_DIR)
DATASET_DIR = os.path.join(ROOT_DIR, "Dataset")
IMAGE_DIR   = os.path.join(ROOT_DIR, "images")

def _find(fname):
    for d in [DATASET_DIR, ROOT_DIR, BASE_DIR]:
        p = os.path.join(d, fname)
        if os.path.exists(p): return p
    return os.path.join(DATASET_DIR, fname)

MASTER_PATH = _find("player_master.csv")
STATS_PATH  = _find("player_stats2.xlsx")

# ── BRAND PALETTE ─────────────────────────────────────────────
NAVY  = "#1B3A6B"
AMBER = "#E8A020"
STEEL = "#4A6FA5"
GREEN = "#1E7A34"
RED   = "#C0392B"
BORDER= "#dde3ee"

# Consistent opponent color mapping
OPP_PALETTE = {
    "Australia":    "#F5A623",
    "West Indies":  "#6B3A2A",
    "South Africa": "#007A4D",
    "New Zealand":  "#7F8C8D",
    "England":      "#003087",
    "Bangladesh":   "#E74C3C",
    "Zimbabwe":     "#C0392B",
    "Pakistan":     "#1A5C2A",
    "Afghanistan":  "#00205B",
    "Netherlands":  "#E55812",
    "Sri Lanka":    "#0066CC",
    "India":        "#FF6600",
    "default":      NAVY,
}

def opp_clr(name):
    return OPP_PALETTE.get(str(name).strip(), OPP_PALETTE["default"])

def opp_clrs(series):
    return [opp_clr(n) for n in series]

# ── GLOBAL CSS ────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&family=Poppins:wght@600;700;800&display=swap');

html,body,[class*="css"]{font-family:'Inter',sans-serif;background:#fff !important;}
#MainMenu,footer,header{visibility:hidden;}
section[data-testid="stSidebar"]{display:none !important;}

.block-container,
[data-testid="stAppViewBlockContainer"],
[data-testid="block-container"]{
    padding:0 60px 60px 60px !important;
    max-width:1560px !important;
    margin:0 auto !important;
    background:#fff !important;
}
[data-testid="stAppViewBlockContainer"]>div:first-child{padding-top:0 !important;}

/* ─ NAVBAR ─ */
.top-header{
    background:#fff; border-bottom:2px solid #1B3A6B;
    padding:0 60px; height:64px;
    display:flex; align-items:center; justify-content:space-between;
    position:fixed;width:100%; top:0; z-index:999;
    box-shadow:0 2px 12px rgba(27,58,107,.08);
    margin:0 -60px 0 -60px;
}
.header-brand{font-family:'Poppins',sans-serif;font-size:1.2rem;font-weight:700;color:#1B3A6B;}
.header-brand span{color:#E8A020;}
.header-nav{display:flex;gap:8px;align-items:center;}
.h-btn{background:#1B3A6B;color:#fff !important;border:none;padding:7px 18px;border-radius:6px;
       font-size:.82rem;font-weight:600;text-decoration:none !important;display:inline-block;transition:background .2s;}
.h-btn:hover{background:#243f7a;}
.h-btn-ghost{background:transparent;color:#1B3A6B !important;border:1.5px solid #1B3A6B;
             padding:6px 16px;border-radius:6px;font-size:.82rem;font-weight:600;
             text-decoration:none !important;display:inline-block;transition:background .2s;}
.h-btn-ghost:hover{background:#f0f4fb;}

/* ─ PAGE TITLE ─ */
.page-title-bar{background:#f7f9fc;border-left:4px solid #1B3A6B;
    padding:14px 60px;margin:40px -60px 28px -60px;}
.page-title-bar h2{font-family:'Poppins',sans-serif;font-size:1.15rem;font-weight:700;color:#1B3A6B;margin:0;}
.page-title-bar p{color:#6B7A8D;font-size:.8rem;margin:3px 0 0 0;}

/* ─ FILTER PANEL ─ */
.filter-wrap{background:#f7f9fc;border:1px solid #dde3ee;border-radius:10px;
             padding:22px 28px;margin-bottom:26px;}
.filter-label{font-size:.75rem;font-weight:600;color:#1B3A6B;
              letter-spacing:.8px;text-transform:uppercase;margin-bottom:6px;}
.scenario-tag{display:inline-block;background:#EBF3FF;color:#1B3A6B;border:1px solid #c0d0eb;
              padding:3px 12px;border-radius:20px;font-size:.72rem;font-weight:700;
              letter-spacing:.4px;margin-top:10px;}
.filter-note{font-size:.72rem;color:#6B7A8D;margin-top:8px;}

/* Streamlit widget overrides */
.stSelectbox>div>div{background:#fff !important;border:1.5px solid #c8d0e0 !important;
    border-radius:7px !important;color:#1B3A6B !important;font-size:.88rem !important;}
.stSelectbox>div>div:focus-within{border-color:#1B3A6B !important;
    box-shadow:0 0 0 3px rgba(27,58,107,.1) !important;}
div[data-testid="stButton"]>button{background:#1B3A6B !important;color:#fff !important;
    font-family:'Poppins',sans-serif !important;font-weight:700 !important;
    font-size:.9rem !important;letter-spacing:1px !important;border:none !important;
    border-radius:7px !important;padding:12px 0 !important;width:100% !important;}
div[data-testid="stButton"]>button:hover{background:#243f7a !important;}

/* ─ SECTION TITLE ─ */
.sec-title{font-family:'Poppins',sans-serif;font-size:1rem;font-weight:700;color:#1B3A6B;
           margin:0 0 14px 0;padding-bottom:8px;border-bottom:2px solid #e8edf5;
           display:flex;align-items:center;gap:8px;}
.sec-title .dot{width:8px;height:8px;background:#1B3A6B;border-radius:50%;display:inline-block;}

/* ─ PROFILE CARD ─ */
.profile-card{background:#fff;border:1px solid #dde3ee;border-radius:12px;
              padding:28px 22px;box-shadow:0 2px 12px rgba(27,58,107,.07);
              position:relative;overflow:hidden;}
.profile-card::before{content:'';position:absolute;top:0;left:0;right:0;height:4px;
    background:linear-gradient(90deg,#1B3A6B 0%,#E8A020 100%);}
.player-name{font-family:'Poppins',sans-serif;font-size:1.2rem;font-weight:700;
             color:#1B3A6B;text-align:center;margin:12px 0 4px 0;}
.player-role-badge{display:block;width:fit-content;margin:0 auto 18px auto;
    padding:3px 14px;border-radius:20px;font-size:.72rem;font-weight:700;
    letter-spacing:.5px;text-transform:uppercase;}
.badge-batsman{background:#EBF3FF;color:#1B3A6B;border:1px solid #c0d0eb;}
.badge-bowler{background:#FFF0E8;color:#C0501A;border:1px solid #f0c0a0;}
.badge-allrounder{background:#F0F8EE;color:#2A7A2A;border:1px solid #a8d4a0;}
.info-row{display:flex;justify-content:space-between;align-items:center;
          padding:9px 0;border-bottom:1px solid #f0f3f8;font-size:.83rem;}
.info-row:last-child{border-bottom:none;}
.info-lbl{color:#6B7A8D;font-weight:500;}
.info-val{color:#1B3A6B;font-weight:600;}

/* ─ BANNER ─ */
.opp-banner{background:linear-gradient(135deg,#1B3A6B,#243f7a);border-radius:10px;
    padding:16px 22px;display:flex;align-items:center;justify-content:space-between;margin-bottom:20px;}
.opp-banner-lbl{color:rgba(255,255,255,.65);font-size:.74rem;text-transform:uppercase;letter-spacing:1px;}
.opp-banner-val{font-family:'Poppins',sans-serif;font-size:1.3rem;font-weight:700;color:#fff;margin-top:2px;}
.opp-banner-sub{color:rgba(255,255,255,.6);font-size:.78rem;margin-top:4px;}
.opp-banner-pct{font-family:'Poppins',sans-serif;font-size:1.6rem;font-weight:700;color:#E8A020;}
.opp-banner-pct-lbl{color:rgba(255,255,255,.65);font-size:.72rem;text-transform:uppercase;letter-spacing:.8px;}

/* ─ INSIGHT CARDS ─ */
.insight-card{background:#f7f9fc;border:1px solid #dde3ee;border-left:4px solid #E8A020;
    border-radius:8px;padding:12px 16px;display:flex;align-items:flex-start;
    gap:12px;margin-bottom:9px;}
.insight-icon{font-size:1.1rem;flex-shrink:0;margin-top:2px;}
.insight-text{font-size:.84rem;color:#2c3e50;line-height:1.5;}
.insight-text strong{color:#1B3A6B;}

/* ─ KPI CARDS ─ */
.kpi-card{background:#fff;border:1px solid #dde3ee;border-radius:10px;
    padding:18px 14px 14px 14px;text-align:center;
    box-shadow:0 2px 8px rgba(27,58,107,.05);
    transition:box-shadow .2s,transform .2s;position:relative;overflow:hidden;}
.kpi-card:hover{box-shadow:0 6px 18px rgba(27,58,107,.12);transform:translateY(-2px);}
.kpi-card::after{content:'';position:absolute;bottom:0;left:0;right:0;height:3px;background:#1B3A6B;}
.kpi-card.hl::after{background:#E8A020;}
.kpi-card.hl-g::after{background:#1E7A34;}
.kpi-val{font-family:'Poppins',sans-serif;font-size:1.65rem;font-weight:700;
         color:#1B3A6B;line-height:1;margin-bottom:6px;}
.kpi-lbl{font-size:.7rem;color:#8899aa;font-weight:600;letter-spacing:.6px;text-transform:uppercase;}
.kpi-sub{font-size:.68rem;color:#aab0bc;margin-top:3px;}

/* ─ CHART CARD ─ */
.chart-card{background:#fff;border:1px solid #dde3ee;border-radius:12px;
    padding:20px 18px 10px 18px;box-shadow:0 2px 10px rgba(27,58,107,.06);margin-bottom:20px;}

/* ─ SCENARIO BANNER ─ */
.sc-banner{background:#f7f9fc;border:1px solid #dde3ee;border-radius:8px;
    padding:13px 20px;margin-bottom:22px;display:flex;align-items:center;gap:14px;}
.sc-pill{background:#1B3A6B;color:#fff;padding:4px 14px;border-radius:20px;
    font-size:.72rem;font-weight:700;letter-spacing:.5px;white-space:nowrap;}
.sc-desc{font-size:.83rem;color:#6B7A8D;line-height:1.4;}

/* ─ DIVIDER / NO DATA ─ */
.divider{border:none;height:1px;background:#e8edf5;margin:28px 0;}
.no-data{text-align:center;padding:50px;color:#8899aa;font-size:.9rem;}
.no-data .icon{font-size:2.5rem;display:block;margin-bottom:10px;}

/* ─ BADGE ─ */
.badge{padding:3px 10px;border-radius:20px;font-size:.7rem;font-weight:700;}
.b-won{background:#EAF6EE;color:#1E7E34;border:1px solid #b8dfc4;}
.b-lose{background:#FDE8E8;color:#C0392B;border:1px solid #f0b8b8;}
.b-nr{background:#F4F4F4;color:#666;border:1px solid #ddd;}

/* ─ ACTION BAR + FOOTER ─ */
.action-bar{background:#f7f9fc;border-top:1px solid #dde3ee;
    padding:26px 60px;margin:32px -60px 0 -60px;
    display:flex;align-items:center;justify-content:space-between;}
.act-btn{display:inline-flex;align-items:center;gap:8px;padding:11px 28px;border-radius:7px;
    font-family:'Poppins',sans-serif;font-weight:700;font-size:.88rem;
    letter-spacing:.5px;cursor:pointer;text-decoration:none !important;border:none;}
.act-report{background:#fff;color:#1B3A6B !important;border:2px solid #1B3A6B;}
.act-report:hover{background:#EBF3FF;}
.act-compare{background:#1B3A6B;color:#fff !important;}
.act-compare:hover{background:#243f7a;}
.footer{background:#1B3A6B;color:rgba(255,255,255,.7);text-align:center;
    padding:18px 60px;font-size:.8rem;margin:0 -60px -60px -60px;letter-spacing:.3px;}
.footer span{color:#E8A020;font-weight:600;}
</style>
""", unsafe_allow_html=True)

# ── DATA LOADING ──────────────────────────────────────────────
@st.cache_data
def load_data():
    master = pd.read_csv(MASTER_PATH)
    master.columns = master.columns.str.strip()
    wb   = openpyxl.load_workbook(STATS_PATH)
    skip = {"Match_id","Opponent","Venue","Country","Match_result",
            "Dismissed_method","Dismissed_Length","Dismissed_against",
            "Dismissed_by","Dismissed _by","Match_date"}
    stats = {}
    for sheet in wb.sheetnames:
        ws   = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            stats[sheet.strip()] = pd.DataFrame(); continue
        cols = [str(c).strip() if c else f"_c{i}" for i, c in enumerate(rows[0])]
        df   = pd.DataFrame(rows[1:], columns=cols)
        df["Match_date"] = pd.to_datetime(df["Match_date"], errors="coerce")
        for c in df.columns:
            if c not in skip:
                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
        stats[sheet.strip()] = df
    return master, stats

try:
    master_df, stats_dict = load_data()
except Exception as e:
    st.error(f" Data load error: {e}"); st.stop()

# ── UTILITY HELPERS ───────────────────────────────────────────
def get_img_b64(filename):
    for d in [IMAGE_DIR, os.path.join(ROOT_DIR, "images"), BASE_DIR]:
        p = os.path.join(d, filename)
        if os.path.exists(p):
            ext  = filename.rsplit(".", 1)[-1].lower()
            mime = "jpeg" if ext in ["jpg","jpeg"] else ext
            with open(p, "rb") as f:
                return f"data:image/{mime};base64,{base64.b64encode(f.read()).decode()}"
    return None

def initials_avatar(name):
    ini = "".join(p[0].upper() for p in name.split()[:2])
    return (f'<div style="width:100px;height:100px;border-radius:50%;background:#EBF3FF;'
            f'border:3px solid #1B3A6B;display:flex;align-items:center;justify-content:center;'
            f'margin:0 auto;font-family:Poppins,sans-serif;font-size:1.8rem;font-weight:700;'
            f'color:#1B3A6B;">{ini}</div>')

def calc_avg(runs_s, not_out_s, innings):
    tot = float(runs_s.sum())
    dis = max(int(innings) - int(not_out_s.sum()), 0)
    return round(tot / dis, 2) if dis > 0 else round(tot, 2)

def pchart(fig, h=310, title=""):
    """Apply consistent chart styling."""
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="#fff",
        font=dict(family="Inter", color="#2c3e50", size=11),
        height=h, margin=dict(l=10, r=10, t=44, b=10),
        title=dict(text=title,
                   font=dict(family="Poppins", size=13, color=NAVY),
                   x=0.02, xanchor="left"),
        legend=dict(bgcolor="rgba(255,255,255,.9)", bordercolor=BORDER,
                    borderwidth=1, font=dict(size=10)),
    )
    fig.update_xaxes(gridcolor="#f0f3f8", linecolor=BORDER,
                     tickfont=dict(size=9), zeroline=False)
    fig.update_yaxes(gridcolor="#f0f3f8", linecolor=BORDER,
                     tickfont=dict(size=9), zeroline=False)
    return fig

def kcard(val, lbl, sub="", v=""):
    return (f'<div class="kpi-card {v}"><div class="kpi-val">{val}</div>'
            f'<div class="kpi-lbl">{lbl}</div>'
            + (f'<div class="kpi-sub">{sub}</div>' if sub else "") + '</div>')

def insight(icon, txt):
    return (f'<div class="insight-card"><span class="insight-icon">{icon}</span>'
            f'<span class="insight-text">{txt}</span></div>')

def get_col(df, *names):
    """Return first existing column from alternatives."""
    for n in names:
        if n in df.columns: return df[n]
    return pd.Series(0, index=df.index)

# ── NAVBAR ─────────────────────────────────────────────────────
st.markdown("""
<div class="top-header">
  <div class="header-brand"> Cricket Player<span> Performance Analysis</span></div>
  <div class="header-nav">
    <a class="h-btn" href="/" target="_self"> Back</a>
    <a class="h-btn" href="/userhome" target="_self"> Home</a>
    <a class="h-btn" href="/admin" target="_self"> Admin</a>
  </div>
</div>
""", unsafe_allow_html=True)

st.markdown("""
<div class="page-title-bar">
  <h2>Opponent Wise Performance Analysis</h2>
  <p>Head-to-head breakdowns, trend analysis, and match-level insights — automatically adapts to your selection</p>
</div>
""", unsafe_allow_html=True)

# ── FILTER PANEL ───────────────────────────────────────────────
fc1, fc2, fc3 = st.columns([3, 3, 2])

with fc1:
    st.markdown('<div class="filter-label"> Select Player</div>', unsafe_allow_html=True)
    sel_player = st.selectbox("Player", sorted(stats_dict.keys()), label_visibility="collapsed")

_raw = stats_dict.get(sel_player, pd.DataFrame())

with fc2:
    st.markdown('<div class="filter-label"> Select Opponent</div>', unsafe_allow_html=True)

    odi_nations = [
        "Afghanistan","Australia","Bangladesh","England","Ireland",
        "New Zealand","Pakistan","South Africa","Sri Lanka","West Indies",
        "Zimbabwe","Netherlands","Scotland","Nepal","Oman",
        "United Arab Emirates","United States","Namibia","Papua New Guinea"
    ]

    sel_opp = st.selectbox(
        "Opponent",
        ["All"] + sorted(odi_nations),
        label_visibility="collapsed"
    )

with fc3:
    st.markdown('<div class="filter-label"> Select Year</div>', unsafe_allow_html=True)

    _yrs = (
        ["All"] + sorted([int(y) for y in _raw["Match_date"].dropna().dt.year.unique()])
        if not _raw.empty else ["All"]
    )

    sel_year = st.selectbox("Year", _yrs, label_visibility="collapsed")

st.markdown('<div class="filter-label">&nbsp;</div>', unsafe_allow_html=True)
analyze = st.button("ANALYZE NOW", use_container_width=True)

# ── SESSION STATE ──────────────────────────────────────────────
if analyze:
    st.session_state.update(ow_ok=True, ow_player=sel_player,
                            ow_opp=sel_opp, ow_year=sel_year)

if not st.session_state.get("ow_ok"):
    st.markdown('<div class="no-data"><span class="icon"></span>'
                'Select player, opponent &amp; year, then click <b>Analyze Now</b></div>',
                unsafe_allow_html=True)
    st.stop()

# Restore from session
sel_player = st.session_state["ow_player"]
sel_opp    = st.session_state["ow_opp"]
sel_year   = st.session_state["ow_year"]
df_raw     = stats_dict.get(sel_player, pd.DataFrame())

# Final scenario after session restore
if sel_opp == "All":
    scenario = 1
elif sel_year == "All":
    scenario = 2
else:
    scenario = 3

SC_LABELS = {
    1: (" Scenario 1 — All-Opponent Comparison",
        "Comparing performance across every opponent"),
    2: (" Scenario 2 — Year-wise vs Opponent",
        f"Year-on-year performance trends vs {sel_opp}"),
    3: (" Scenario 3 — Match Level Analysis",
        f"Match-by-match breakdown vs {sel_opp} in {sel_year}"),
}

# ── PLAYER META ────────────────────────────────────────────────
prow = master_df[master_df["PlayerName"].str.strip() == sel_player]
if not prow.empty:
    pi        = prow.iloc[0]
    p_country = pi.get("Country",     "India")
    p_role    = pi.get("Role",        "Batsman")
    p_bat     = pi.get("BattingStyle","RHB")
    p_bowl    = pi.get("BowlingStyle","-")
    p_dob     = pi.get("DOB",         "N/A")
    p_debut   = pi.get("DebutYear",   "N/A")
    p_img     = pi.get("ImageFileName","")
else:
    p_country="India"; p_role="Batsman"; p_bat="RHB"
    p_bowl="-"; p_dob="N/A"; p_debut="N/A"; p_img=""

rl     = p_role.lower()
is_bat = "batsman" in rl
is_bowl= "bowler"  in rl
is_all = "all-rounder" in rl or "all rounder" in rl
bcls   = ("badge-bowler" if is_bowl else
          ("badge-allrounder" if is_all else "badge-batsman"))

# ── BUILD WORKING DATAFRAME ────────────────────────────────────
if df_raw.empty:
    st.markdown('<div class="no-data"><span class="icon">📭</span>'
                'No data for this player.</div>', unsafe_allow_html=True)
    st.stop()

wdf = df_raw.copy()
if sel_opp != "All":
    wdf = wdf[wdf["Opponent"] == sel_opp]
if sel_year != "All":
    wdf = wdf[wdf["Match_date"].dt.year == int(sel_year)]
wdf = wdf.sort_values("Match_date")

def bat_rows(df):
    return df[df["Batted_Flag"] == 1] if "Batted_Flag" in df.columns else df

def bowl_rows(df):
    return df[df["Bowl_Flag"] == 1] if "Bowl_Flag" in df.columns else df

bat_df  = bat_rows(wdf)
bowl_df = bowl_rows(wdf)

if bat_df.empty and bowl_df.empty:
    st.markdown('<div class="no-data"><span class="icon"></span>'
                'No data for the selected filters.</div>', unsafe_allow_html=True)
    st.stop()

total_m   = len(wdf)
opp_label = sel_opp if sel_opp != "All" else "All Opponents"
wins_m    = int((wdf["Match_result"] == "Won").sum())
losses_m  = int((wdf["Match_result"] == "Lose").sum())
win_pct   = round(wins_m / max(total_m, 1) * 100, 1)

# ══════════════════════════════════════════════════════════════
#  SECTION 1 — PROFILE + KEY INSIGHTS
# ══════════════════════════════════════════════════════════════
col_p, col_i = st.columns([1, 2.2], gap="large")

with col_p:
    st.markdown('<div class="sec-title"><span class="dot"></span>Player Profile</div>',
                unsafe_allow_html=True)
    img_src  = get_img_b64(p_img) if p_img else None
    img_html = (f'<img src="{img_src}" style="width:100px;height:100px;border-radius:50%;'
                f'object-fit:cover;border:3px solid #1B3A6B;display:block;margin:0 auto;">'
                if img_src else initials_avatar(sel_player))
    st.markdown(f"""
    <div class="profile-card">
      {img_html}
      <div class="player-name">{sel_player}</div>
      <div class="player-role-badge {bcls}">{p_role}</div>
      <div class="info-row"><span class="info-lbl">Batting</span>
        <span class="info-val">{p_bat}</span></div>
      <div class="info-row"><span class="info-lbl">Bowling</span>
        <span class="info-val">{p_bowl}</span></div>
      <div class="info-row"><span class="info-lbl">Date of Birth</span>
        <span class="info-val">{p_dob}</span></div>
      <div class="info-row"><span class="info-lbl">Debut Year</span>
        <span class="info-val">{p_debut}</span></div>
      <div class="info-row"><span class="info-lbl">vs</span>
        <span class="info-val">{opp_label}</span></div>
    </div>""", unsafe_allow_html=True)

with col_i:
    # Summary banner
    nr_m = total_m - wins_m - losses_m
    st.markdown(f"""
    <div class="opp-banner">
      <div>
        <div class="opp-banner-lbl">
          {'Head-to-Head' if sel_opp != 'All' else 'All-Opponent Overview'}
        </div>
        <div class="opp-banner-val">
          {'vs ' + sel_opp if sel_opp != 'All' else 'All Opponents'}
        </div>
        <div class="opp-banner-sub">
          {wins_m} Won &nbsp;·&nbsp; {losses_m} Lost &nbsp;·&nbsp;
          {nr_m} N/R &nbsp;·&nbsp; {total_m} Matches
        </div>
      </div>
      <div style="text-align:right;">
        <div class="opp-banner-pct">{win_pct}%</div>
        <div class="opp-banner-pct-lbl">Win Rate</div>
      </div>
    </div>""", unsafe_allow_html=True)

    st.markdown('<div class="sec-title"><span class="dot"></span>Key Insights</div>',
                unsafe_allow_html=True)

    # ── Compute insights from FULL dataset (not filtered) ────
    full_bat  = bat_rows(df_raw)
    full_bowl = bowl_rows(df_raw)
    ins_html  = ""

    if is_bat and not is_all:
        # Best / weak opponent by average
        if not full_bat.empty and full_bat["Opponent"].nunique() > 1:
            oa = (full_bat.groupby("Opponent")
                  .apply(lambda g: calc_avg(g["Runs"],
                         g.get("Not_out", pd.Series(0, index=g.index)), len(g)))
                  .reset_index(name="Avg"))
            best_r = oa.loc[oa["Avg"].idxmax()]
            weak_r = oa.loc[oa["Avg"].idxmin()]
            ins_html += insight("",
                f"Best opponent: <strong>{best_r['Opponent']}</strong> "
                f"(Avg <strong>{best_r['Avg']}</strong>)")
            ins_html += insight("",
                f"Weakest opponent: <strong>{weak_r['Opponent']}</strong> "
                f"(Avg <strong>{weak_r['Avg']}</strong>)")
        # Best SR opponent
        if not full_bat.empty and full_bat["Opponent"].nunique() > 1:
            sr_m = full_bat.groupby("Opponent")["Strike_Rate"].mean().round(1)
            bsr  = sr_m.idxmax()
            ins_html += insight("",
                f"Highest strike rate vs <strong>{bsr}</strong> "
                f"(SR <strong>{sr_m[bsr]}</strong>)")
        # Current scope
        if not bat_df.empty:
            tr   = int(bat_df["Runs"].sum())
            bnd  = int(bat_df["4s"].sum() * 4 + bat_df["6s"].sum() * 6)
            bpct = round(bnd / max(tr, 1) * 100, 1)
            ins_html += insight("",
                f"<strong>{bpct}%</strong> of runs from boundaries "
                f"({bnd} of {tr} runs vs {opp_label})")
            h50 = bat_df[bat_df["Runs"] >= 50]
            if not h50.empty:
                wp50 = round((h50["Match_result"] == "Won").sum() / len(h50) * 100, 1)
                ins_html += insight("",
                    f"Team wins <strong>{wp50}%</strong> when {sel_player} "
                    f"scores 50+ vs {opp_label}")

    elif is_bowl and not is_all:
        if not full_bowl.empty and full_bowl["Opponent"].nunique() > 1:
            wk_m = full_bowl.groupby("Opponent")["Wickets"].sum().reset_index()
            bw   = wk_m.loc[wk_m["Wickets"].idxmax()]
            wkw  = wk_m.loc[wk_m["Wickets"].idxmin()]
            ins_html += insight("",
                f"Most wickets vs <strong>{bw['Opponent']}</strong> "
                f"(<strong>{int(bw['Wickets'])}</strong> wickets)")
            ins_html += insight("",
                f"Fewest wickets vs <strong>{wkw['Opponent']}</strong> "
                f"(<strong>{int(wkw['Wickets'])}</strong> wickets)")
            eco_m = full_bowl.groupby("Opponent")["Economy"].mean().round(2)
            be    = eco_m.idxmin()
            ins_html += insight("",
                f"Best economy vs <strong>{be}</strong> "
                f"(<strong>{eco_m[be]:.2f}</strong> RPO)")
        if not bowl_df.empty:
            tb  = int(bowl_df["Balls_Bowled"].sum()) if "Balls_Bowled" in bowl_df else 0
            db  = int(bowl_df.get("Dot_Balls", pd.Series(0)).sum())
            ins_html += insight("",
                f"<strong>{round(db/max(tb,1)*100,1)}%</strong> dot ball rate "
                f"vs {opp_label} ({db} dots from {tb} balls)")
            tw  = int(bowl_df["Wickets"].sum())
            h3  = int((bowl_df["Wickets"] >= 3).sum())
            ins_html += insight("",
                f"<strong>{tw}</strong> wickets with "
                f"<strong>{h3}</strong> three-wicket+ hauls vs {opp_label}")

    else:  # All-rounder
        if not full_bat.empty and full_bat["Opponent"].nunique() > 1:
            oa2  = (full_bat.groupby("Opponent")
                    .apply(lambda g: calc_avg(g["Runs"],
                           g.get("Not_out", pd.Series(0, index=g.index)), len(g)))
                    .reset_index(name="Avg"))
            best2 = oa2.loc[oa2["Avg"].idxmax()]
            ins_html += insight("",
                f"Best batting vs <strong>{best2['Opponent']}</strong> "
                f"(Avg <strong>{best2['Avg']}</strong>)")
        if not full_bowl.empty and full_bowl["Opponent"].nunique() > 1:
            wb2  = full_bowl.groupby("Opponent")["Wickets"].sum()
            bo2  = wb2.idxmax()
            ins_html += insight("",
                f"Most bowling impact vs <strong>{bo2}</strong> "
                f"(<strong>{int(wb2[bo2])}</strong> wickets)")
        if not bat_df.empty and not bowl_df.empty:
            tr2 = int(bat_df["Runs"].sum())
            tw2 = int(bowl_df["Wickets"].sum())
            ins_html += insight("",
                f"All-round contribution vs {opp_label}: "
                f"<strong>{tr2}</strong> runs + <strong>{tw2}</strong> wickets")
        if not bat_df.empty:
            h30 = bat_df[bat_df["Runs"] >= 30]
            if not h30.empty:
                wp3 = round((h30["Match_result"] == "Won").sum() / len(h30) * 100, 1)
                ins_html += insight("",
                    f"Team wins <strong>{wp3}%</strong> when {sel_player} "
                    f"scores 30+ vs {opp_label}")

    st.markdown(ins_html or
        '<div class="no-data" style="padding:20px">No insights available.</div>',
        unsafe_allow_html=True)

# ══════════════════════════════════════════════════════════════
#  SECTION 2 — KPI CARDS
# ══════════════════════════════════════════════════════════════
st.markdown('<hr class="divider">', unsafe_allow_html=True)
st.markdown('<div class="sec-title"><span class="dot"></span>'
            'Key Performance Indicators</div>', unsafe_allow_html=True)

if is_bat and not is_all:
    inn    = len(bat_df)
    tot_r  = int(bat_df["Runs"].sum())
    no_s   = bat_df.get("Not_out", pd.Series(0, index=bat_df.index))
    avg    = calc_avg(bat_df["Runs"], no_s, inn)
    avg_sr = round(bat_df["Strike_Rate"].replace(0, float("nan")).mean(), 1)
    f4     = int(bat_df["4s"].sum()); f6 = int(bat_df["6s"].sum())
    fif    = int(bat_df["50s"].sum()) if "50s" in bat_df else 0
    huns   = int(bat_df["100s"].sum()) if "100s" in bat_df else 0
    if not bat_df.empty:
        hs_r = bat_df.loc[bat_df["Runs"].idxmax()]
        hs   = f"{int(hs_r['Runs'])}{'*' if int(hs_r.get('Not_out',0)) else ''}"
    else:
        hs = "0"
    kpis = [
        (str(inn),          "Innings",    f"vs {opp_label}", ""),
        (str(tot_r),        "Total Runs", "Aggregate",       "hl"),
        (str(avg),          "Average",    "Runs/dismissal",  ""),
        (str(avg_sr),       "Strike Rate","Avg SR",          ""),
        (hs,                "Best Score", "Highest innings", "hl"),
        (f"{f4}/{f6}",      "4s / 6s",   "Boundaries",      ""),
        (f"{fif}/{huns}",   "50s / 100s","Milestones",       "hl-g"),
    ]

elif is_bowl and not is_all:
    inn     = len(bowl_df)
    t_balls = int(bowl_df["Balls_Bowled"].sum()) if "Balls_Bowled" in bowl_df else 0
    ovs_s   = f"{t_balls//6}.{t_balls%6}"
    t_wkts  = int(bowl_df["Wickets"].sum()) if "Wickets" in bowl_df else 0
    rc      = int(bowl_df["Runs_Conceded"].sum()) if "Runs_Conceded" in bowl_df else 0
    bowl_a  = round(rc / max(t_wkts, 1), 2)
    avg_eco = round(bowl_df["Economy"].replace(0, float("nan")).mean(), 2) if "Economy" in bowl_df else 0
    dots_b  = int(get_col(bowl_df, "Dot_Balls").sum())
    if t_wkts > 0:
        bf_row = bowl_df.sort_values(["Wickets","Runs_Conceded"],
                                     ascending=[False,True]).iloc[0]
        bfig   = f"{int(bf_row['Wickets'])}/{int(bf_row.get('Runs_Conceded',0))}"
    else:
        bfig = "0/0"
    kpis = [
        (str(inn),    "Matches",     f"vs {opp_label}", ""),
        (ovs_s,       "Overs",       "Total bowled",    ""),
        (str(t_wkts), "Wickets",     "Total",           "hl"),
        (str(bowl_a), "Bowl Average","Runs/wicket",     ""),
        (str(dots_b), "Dot Balls",   "Scoreless",       ""),
        (str(avg_eco),"Economy",     "Runs/over",       ""),
        (bfig,        "Best Figures","Wkts/Runs",       "hl"),
    ]

else:  # All-rounder
    b_inn   = len(bat_df)
    tot_r2  = int(bat_df["Runs"].sum()) if not bat_df.empty else 0
    no_s2   = (bat_df.get("Not_out", pd.Series(0, index=bat_df.index))
               if not bat_df.empty else pd.Series([0]))
    avg2    = calc_avg(bat_df["Runs"] if not bat_df.empty else pd.Series([0]), no_s2, max(b_inn,1))
    avg_sr2 = round(bat_df["Strike_Rate"].replace(0,float("nan")).mean(),1) if not bat_df.empty else 0
    fif2    = int(bat_df["50s"].sum()) if not bat_df.empty and "50s" in bat_df else 0
    huns2   = int(bat_df["100s"].sum()) if not bat_df.empty and "100s" in bat_df else 0
    t_wkts2 = int(bowl_df["Wickets"].sum()) if not bowl_df.empty else 0
    avg_eco2= round(bowl_df["Economy"].replace(0,float("nan")).mean(),2) if not bowl_df.empty and "Economy" in bowl_df else 0
    t_ball2 = int(bowl_df["Balls_Bowled"].sum()) if not bowl_df.empty and "Balls_Bowled" in bowl_df else 0
    ovs_s2  = f"{t_ball2//6}.{t_ball2%6}"
    kpis = [
        (str(b_inn),       "Innings",    f"vs {opp_label}", ""),
        (str(tot_r2),      "Total Runs", "Batting",         "hl"),
        (str(avg2),        "Bat Average","Runs/dismissal",  ""),
        (str(avg_sr2),     "Strike Rate","Avg SR",          ""),
        (f"{fif2}/{huns2}","50s / 100s", "Milestones",      "hl-g"),
        (ovs_s2,           "Overs",      "Bowled",          ""),
        (str(t_wkts2),     "Wickets",    "Bowling",         "hl"),
        (str(avg_eco2),    "Economy",    "Runs/over",       ""),
    ]

kcols = st.columns(len(kpis))
for i, (v, l, s, var) in enumerate(kpis):
    with kcols[i]:
        st.markdown(kcard(v, l, s, var), unsafe_allow_html=True)
st.divider()

# ══════════════════════════════════════════════════════════════
#  SECTION 3 — CHARTS 
# ══════════════════════════════════════════════════════════════

# ── AGGREGATE HELPERS ─────────────────────────────────────────
def agg_bat(df, group_col="Opponent"):
    out = []
    for key, g in df.groupby(group_col):
        inn = len(g)
        tr  = int(g["Runs"].sum())
        no  = g.get("Not_out", pd.Series(0, index=g.index))
        avg_= calc_avg(g["Runs"], no, inn)
        sr_ = round(g["Strike_Rate"].replace(0,float("nan")).mean(), 1)
        f_  = int(g["4s"].sum()); s_ = int(g["6s"].sum())
        d_  = int(g["Dots"].sum())
        sg_ = int(g["1s_2s_3s"].sum()) if "1s_2s_3s" in g else 0
        fi_ = int(g["50s"].sum()) if "50s" in g else 0
        hu_ = int(g["100s"].sum()) if "100s" in g else 0
        dm_ = g[~g["Dismissed_method"].astype(str).isin(
              ["Not Applicable","0","nan",""])]
        bnd_r = f_*4 + s_*6
        out.append({group_col:key, "Innings":inn, "Runs":tr, "Average":avg_,
                    "SR":sr_, "Fours":f_, "Sixes":s_, "Dots":d_, "Singles":sg_,
                    "Fifties":fi_, "Hundreds":hu_, "Bnd_Pct":round(bnd_r/max(tr,1)*100,1)})
    return pd.DataFrame(out)

def agg_bowl(df, group_col="Opponent"):
    out = []
    for key, g in df.groupby(group_col):
        wk = int(g["Wickets"].sum())
        rc = int(g["Runs_Conceded"].sum()) if "Runs_Conceded" in g else 0
        ec = round(g["Economy"].replace(0,float("nan")).mean(), 2) if "Economy" in g else 0
        bl = int(g["Balls_Bowled"].sum()) if "Balls_Bowled" in g else 0
        bsr= round(bl/max(wk,1),1)
        ba = round(rc/max(wk,1),2)
        fd = int(g["Full_Dismissals"].sum()) if "Full_Dismissals" in g else 0
        gd = int(get_col(g,"Good_Dissmisals","Good_Dismissals").sum())
        sd = int(get_col(g,"Short_Dissmisals","Short_Dismissals").sum())
        rw = int(g["Wicket_RHB"].sum()) if "Wicket_RHB" in g else 0
        lw = int(g["Wicket_LHB"].sum()) if "Wicket_LHB" in g else 0
        out.append({group_col:key, "Wickets":wk, "Runs_C":rc, "Economy":ec,
                    "Bowl_SR":bsr, "Bowl_Avg":ba, "Full":fd, "Good":gd,
                    "Short":sd, "RHB_W":rw, "LHB_W":lw})
    return pd.DataFrame(out)

# ══════════════════════════════════════════════════════════════
#  SCENARIO 1 — ALL OPPONENTS COMPARISON
# ══════════════════════════════════════════════════════════════
if scenario == 1:
    base = df_raw.copy()
    if sel_year != "All":
        base = base[base["Match_date"].dt.year == int(sel_year)]

    # ─── BATSMAN ──────────────────────────────────────────────
    if is_bat and not is_all:
        ogg = agg_bat(bat_rows(base))
        if ogg.empty:
            st.markdown('<div class="no-data">No batting data available.</div>',
                        unsafe_allow_html=True)
        else:
            clrs = opp_clrs(ogg["Opponent"])

            # ROW 1 — 3 charts
            r1a, r1b, r1c = st.columns(3, gap="large")

            with r1a:  # 1. Runs vs Opponent — horizontal bar
                fig = go.Figure(go.Bar(
                    y=ogg["Opponent"], x=ogg["Runs"], orientation="h",
                    marker=dict(color=clrs, line=dict(color="white",width=1)),
                    text=ogg["Runs"], textposition="outside",
                    textfont=dict(family="Poppins",size=11),
                ))
                pchart(fig, 300, "Runs vs Each Opponent")
                fig.update_xaxes(title_text="Runs")
                st.plotly_chart(fig, use_container_width=True)

            with r1b:  # 2. Batting Average — column chart
                avg_clr = ["#1E7A34" if a >= 50 else (AMBER if a >= 30 else RED)
                           for a in ogg["Average"]]
                fig2 = go.Figure(go.Bar(
                    x=ogg["Opponent"], y=ogg["Average"],
                    marker=dict(color=avg_clr, line=dict(color="white",width=1)),
                    text=ogg["Average"], textposition="outside",
                    textfont=dict(family="Poppins",size=11),
                ))
                fig2.add_hline(y=50, line_dash="dash", line_color=AMBER,
                    annotation_text="Avg 50",
                    annotation_font=dict(color=AMBER,size=9))
                pchart(fig2, 300, "Batting Average vs Opponent")
                fig2.update_xaxes(tickangle=-20)
                fig2.update_yaxes(title_text="Average")
                st.plotly_chart(fig2, use_container_width=True)

            with r1c:  # 3. Strike Rate — dot plot
                sr_clr = [GREEN if s >= 100 else (AMBER if s >= 70 else RED)
                          for s in ogg["SR"]]
                fig3 = go.Figure(go.Scatter(
                    x=ogg["Opponent"], y=ogg["SR"],
                    mode="markers+text",
                    text=[str(s) for s in ogg["SR"]],
                    textposition="top center",
                    textfont=dict(family="Poppins",size=9,color=NAVY),
                    marker=dict(size=16, color=sr_clr,
                                line=dict(color="white",width=2)),
                ))
                fig3.add_hline(y=100, line_dash="dash", line_color=GREEN,
                    annotation_text="SR 100",
                    annotation_font=dict(color=GREEN,size=9))
                pchart(fig3, 300, "Strike Rate vs Opponent")
                fig3.update_xaxes(tickangle=-20)
                fig3.update_yaxes(title_text="Strike Rate")
                st.plotly_chart(fig3, use_container_width=True)

            # ROW 2 — 2 charts
            r2a, r2b = st.columns(2, gap="large")

            with r2a:  # 4. Boundary % — stacked 100% bar
                fig4 = go.Figure()
                fig4.add_trace(go.Bar(name="4s Runs",  x=ogg["Opponent"],
                    y=ogg["Fours"]*4,   marker_color=NAVY))
                fig4.add_trace(go.Bar(name="6s Runs",  x=ogg["Opponent"],
                    y=ogg["Sixes"]*6,   marker_color=AMBER))
                fig4.add_trace(go.Bar(name="1s/2s/3s", x=ogg["Opponent"],
                    y=ogg["Singles"],   marker_color=STEEL))
                fig4.add_trace(go.Bar(name="Dots",     x=ogg["Opponent"],
                    y=ogg["Dots"],      marker_color="#D5DDE5"))
                fig4.update_layout(barmode="stack", barnorm="percent")
                pchart(fig4, 320, "Boundary % Composition vs Opponent")
                fig4.update_xaxes(tickangle=-20)
                fig4.update_yaxes(title_text="%")
                st.plotly_chart(fig4, use_container_width=True)

            with r2b:  # 5. Dismissal type vs Opponent — heatmap
                b_all = bat_rows(base)
                dis_v = b_all[~b_all["Dismissed_method"].astype(str).isin(
                    ["Not Applicable","0","nan",""])]
                if not dis_v.empty:
                    heat = dis_v.groupby(
                        ["Opponent","Dismissed_method"]).size().unstack(fill_value=0)
                    fig5 = go.Figure(go.Heatmap(
                        z=heat.values, x=heat.columns.tolist(),
                        y=heat.index.tolist(),
                        colorscale=[[0,"#EBF3FF"],[0.5,"#6B9FCC"],[1,NAVY]],
                        text=heat.values, texttemplate="%{text}",
                        textfont=dict(family="Poppins",size=12),
                        showscale=True,
                        colorbar=dict(thickness=12,len=0.8),
                    ))
                    pchart(fig5, 320, "Dismissal Type vs Opponent (Heatmap)")
                    fig5.update_xaxes(title_text="Dismissal Type",tickangle=-15)
                    fig5.update_yaxes(title_text="Opponent")
                    st.plotly_chart(fig5, use_container_width=True)
                else:
                    st.markdown(
                        '<div class="no-data">No dismissal data available.</div>',
                        unsafe_allow_html=True)

    # ─── BOWLER ───────────────────────────────────────────────
    elif is_bowl and not is_all:
        bgg = agg_bowl(bowl_rows(base))
        if bgg.empty:
            st.markdown('<div class="no-data">No bowling data available.</div>',
                        unsafe_allow_html=True)
        else:
            clrs = opp_clrs(bgg["Opponent"])

            r1a, r1b, r1c = st.columns(3, gap="large")

            with r1a:  # 1. Wickets — horizontal bar
                fig = go.Figure(go.Bar(
                    y=bgg["Opponent"], x=bgg["Wickets"], orientation="h",
                    marker=dict(color=clrs, line=dict(color="white",width=1)),
                    text=bgg["Wickets"], textposition="outside",
                    textfont=dict(family="Poppins",size=11),
                ))
                pchart(fig, 300, "Wickets vs Each Opponent")
                fig.update_xaxes(title_text="Wickets")
                st.plotly_chart(fig, use_container_width=True)
                st.markdown('</div>', unsafe_allow_html=True)

            with r1b:  # 2. Economy — dot plot
                eco_clr = [GREEN if e <= 6 else (AMBER if e <= 8 else RED)
                           for e in bgg["Economy"]]
                fig2 = go.Figure(go.Scatter(
                    x=bgg["Opponent"], y=bgg["Economy"],
                    mode="markers+text",
                    text=[str(e) for e in bgg["Economy"]],
                    textposition="top center",
                    textfont=dict(family="Poppins",size=9,color=NAVY),
                    marker=dict(size=16, color=eco_clr,
                                line=dict(color="white",width=2)),
                ))
                fig2.add_hline(y=7, line_dash="dash", line_color=AMBER,
                    annotation_text="Economy 7",
                    annotation_font=dict(color=AMBER,size=9))
                pchart(fig2, 300, "Economy vs Opponent")
                fig2.update_xaxes(tickangle=-20)
                fig2.update_yaxes(title_text="Economy (RPO)")
                st.plotly_chart(fig2, use_container_width=True)

            with r1c:  # 3. Wickets by Length — stacked column
                fig3 = go.Figure()
                fig3.add_trace(go.Bar(name="Good Length", x=bgg["Opponent"],
                    y=bgg["Good"], marker_color=NAVY))
                fig3.add_trace(go.Bar(name="Full Length", x=bgg["Opponent"],
                    y=bgg["Full"], marker_color=AMBER))
                fig3.add_trace(go.Bar(name="Short",       x=bgg["Opponent"],
                    y=bgg["Short"], marker_color=STEEL))
                fig3.update_layout(barmode="stack")
                pchart(fig3, 300, "Wickets by Length vs Opponent")
                fig3.update_xaxes(tickangle=-20)
                fig3.update_yaxes(title_text="Wickets")
                st.plotly_chart(fig3, use_container_width=True)

            r2a, r2b = st.columns(2, gap="large")

            with r2a:  # 4. RHB vs LHB — grouped bar
                fig4 = go.Figure()
                fig4.add_trace(go.Bar(name="vs RHB", x=bgg["Opponent"],
                    y=bgg["RHB_W"], marker_color=NAVY))
                fig4.add_trace(go.Bar(name="vs LHB", x=bgg["Opponent"],
                    y=bgg["LHB_W"], marker_color=AMBER))
                fig4.update_layout(barmode="group")
                pchart(fig4, 320, "Wickets vs RHB / LHB by Opponent")
                fig4.update_xaxes(tickangle=-20)
                fig4.update_yaxes(title_text="Wickets")
                st.plotly_chart(fig4, use_container_width=True)

            with r2b:  # 5. Bowling Average — color intensity bar
                ba_clr = [GREEN if b <= 25 else (AMBER if b <= 35 else RED)
                          for b in bgg["Bowl_Avg"]]
                fig5 = go.Figure(go.Bar(
                    x=bgg["Opponent"], y=bgg["Bowl_Avg"],
                    marker=dict(color=ba_clr, line=dict(color="white",width=1)),
                    text=bgg["Bowl_Avg"], textposition="outside",
                    textfont=dict(family="Poppins",size=11),
                ))
                fig5.add_hline(y=30, line_dash="dash", line_color=AMBER,
                    annotation_text="Avg 30",
                    annotation_font=dict(color=AMBER,size=9))
                pchart(fig5, 320, "Bowling Average vs Opponent")
                fig5.update_xaxes(tickangle=-20)
                fig5.update_yaxes(title_text="Runs per Wicket")
                st.plotly_chart(fig5, use_container_width=True)

    # ─── ALL-ROUNDER ──────────────────────────────────────────
    else:
        ogg = agg_bat(bat_rows(base))
        bgg = agg_bowl(bowl_rows(base))

        r1a, r1b = st.columns(2, gap="large")
        with r1a:  # Bat vs Ball scatter
            if not ogg.empty and not bgg.empty:
                mg = pd.merge(ogg[["Opponent","Runs","Average"]],
                              bgg[["Opponent","Wickets","Economy"]],
                              on="Opponent", how="outer").fillna(0)
                fig = go.Figure(go.Scatter(
                    x=mg["Runs"], y=mg["Wickets"],
                    mode="markers+text", text=mg["Opponent"],
                    textposition="top center",
                    textfont=dict(family="Poppins",size=9),
                    marker=dict(size=14, color=opp_clrs(mg["Opponent"]),
                                line=dict(color="white",width=2)),
                ))
                pchart(fig, 310, "Bat vs Ball Impact (Runs vs Wickets)")
                fig.update_xaxes(title_text="Runs")
                fig.update_yaxes(title_text="Wickets")
                st.plotly_chart(fig, use_container_width=True)

        with r1b:  # Runs horizontal bar
            if not ogg.empty:
                fig2 = go.Figure(go.Bar(
                    y=ogg["Opponent"], x=ogg["Runs"], orientation="h",
                    marker=dict(color=opp_clrs(ogg["Opponent"]),
                                line=dict(color="white",width=1)),
                    text=ogg["Runs"], textposition="outside",
                ))
                pchart(fig2, 310, "Runs vs Each Opponent")
                fig2.update_xaxes(title_text="Runs")
                st.plotly_chart(fig2, use_container_width=True)

        r2a, r2b = st.columns(2, gap="large")
        with r2a:  # Wickets bar
            if not bgg.empty:
                fig3 = go.Figure(go.Bar(
                    x=bgg["Opponent"], y=bgg["Wickets"],
                    marker=dict(color=opp_clrs(bgg["Opponent"]),
                                line=dict(color="white",width=1)),
                    text=bgg["Wickets"], textposition="outside",
                ))
                pchart(fig3, 310, "Wickets vs Each Opponent")
                fig3.update_xaxes(tickangle=-20)
                fig3.update_yaxes(title_text="Wickets")
                st.plotly_chart(fig3, use_container_width=True)

        with r2b:  # Batting average bar
            if not ogg.empty:
                fig4 = go.Figure(go.Bar(
                    x=ogg["Opponent"], y=ogg["Average"],
                    marker=dict(color=opp_clrs(ogg["Opponent"]),
                                line=dict(color="white",width=1)),
                    text=ogg["Average"], textposition="outside",
                ))
                fig4.add_hline(y=50, line_dash="dash", line_color=AMBER,
                    annotation_text="Avg 50",
                    annotation_font=dict(color=AMBER,size=9))
                pchart(fig4, 310, "Batting Average vs Each Opponent")
                fig4.update_xaxes(tickangle=-20)
                fig4.update_yaxes(title_text="Average")
                st.plotly_chart(fig4, use_container_width=True)

# ══════════════════════════════════════════════════════════════
#  SCENARIO 2 — YEAR-WISE TRENDS vs SPECIFIC OPPONENT
# ══════════════════════════════════════════════════════════════
elif scenario == 2:
    base2 = df_raw[df_raw["Opponent"] == sel_opp].copy()
    base2["_yr"] = base2["Match_date"].dt.year.astype(str)
    oc = opp_clr(sel_opp)

    # ─── BATSMAN ──────────────────────────────────────────────
    if is_bat and not is_all:
        b2 = bat_rows(base2)
        if b2.empty:
            st.markdown('<div class="no-data">No year-wise batting data.</div>',
                        unsafe_allow_html=True)
        else:
            yrg = agg_bat(b2, "_yr").sort_values("_yr")
            x   = yrg["_yr"].astype(str)

            r1a, r1b = st.columns(2, gap="large")
            with r1a:
                fig = go.Figure(go.Bar(
                    x=x, y=yrg["Runs"],
                    marker=dict(color=[oc]*len(x), line=dict(color="white",width=1)),
                    text=yrg["Runs"], textposition="outside",
                    textfont=dict(family="Poppins",size=11,color=NAVY),
                ))
                pchart(fig, 300, f"Runs per Year vs {sel_opp}")
                fig.update_yaxes(title_text="Runs")
                fig.update_xaxes(type="category")
                st.plotly_chart(fig, use_container_width=True)

            with r1b:
                fig2 = go.Figure()
                fig2.add_trace(go.Scatter(
                    x=x, y=yrg["Average"], mode="lines+markers+text",
                    name="Average", line=dict(color=NAVY,width=2.5),
                    marker=dict(size=10,color=NAVY,line=dict(color="white",width=2)),
                    text=yrg["Average"], textposition="top center",
                    textfont=dict(family="Poppins",size=9,color=NAVY)))
                fig2.add_trace(go.Scatter(
                    x=x, y=yrg["SR"], mode="lines+markers+text",
                    name="Strike Rate", line=dict(color=AMBER,width=2,dash="dot"),
                    marker=dict(size=8,color=AMBER,line=dict(color="white",width=2)),
                    text=yrg["SR"], textposition="bottom center",
                    textfont=dict(family="Poppins",size=9,color=AMBER)))
                pchart(fig2, 300, f"Average vs Strike Rate — Year-wise vs {sel_opp}")
                st.plotly_chart(fig2, use_container_width=True)

            r2a, r2b = st.columns(2, gap="large")
            with r2a:
                fig3 = go.Figure()
                fig3.add_trace(go.Bar(name="4s Runs", x=x,
                    y=yrg["Fours"]*4, marker_color=NAVY))
                fig3.add_trace(go.Bar(name="6s Runs", x=x,
                    y=yrg["Sixes"]*6, marker_color=AMBER))
                fig3.add_trace(go.Bar(name="1s/2s/3s", x=x,
                    y=yrg["Singles"], marker_color=STEEL))
                fig3.add_trace(go.Bar(name="Dots", x=x,
                    y=yrg["Dots"], marker_color="#D5DDE5"))
                fig3.update_layout(barmode="stack", barnorm="percent")
                pchart(fig3, 300, f"Boundary % per Year vs {sel_opp}")
                fig3.update_yaxes(title_text="%")
                st.plotly_chart(fig3, use_container_width=True)

            with r2b:
                dis2 = b2[~b2["Dismissed_method"].astype(str).isin(
                    ["Not Applicable","0","nan",""])]
                if not dis2.empty:
                    dm2 = dis2.groupby("Dismissed_method").size().reset_index(name="Count")
                    fig4 = go.Figure(go.Pie(
                        labels=dm2["Dismissed_method"], values=dm2["Count"],
                        hole=0.45,
                        marker=dict(colors=[NAVY,AMBER,STEEL,GREEN,RED],
                                    line=dict(color="white",width=2)),
                        textfont=dict(family="Poppins",size=12),
                        textinfo="label+percent"))
                    fig4.update_layout(
                        paper_bgcolor="rgba(0,0,0,0)", height=300,
                        margin=dict(l=10,r=10,t=44,b=10),
                        title=dict(text=f"Dismissal Types vs {sel_opp}",
                            font=dict(family="Poppins",size=13,color=NAVY),x=0.02),
                        showlegend=True,
                        legend=dict(font=dict(size=10),bgcolor="rgba(255,255,255,.9)"))
                    st.plotly_chart(fig4, use_container_width=True)
                else:
                    st.markdown('<div class="no-data">No dismissal data.</div>',
                                unsafe_allow_html=True)

    # ─── BOWLER ───────────────────────────────────────────────
    elif is_bowl and not is_all:
        bw2 = bowl_rows(base2)
        if bw2.empty:
            st.markdown('<div class="no-data">No year-wise bowling data.</div>',
                        unsafe_allow_html=True)
        else:
            yrg2 = agg_bowl(bw2, "_yr").sort_values("_yr")
            x2   = yrg2["_yr"].astype(str)

            r1a, r1b = st.columns(2, gap="large")
            with r1a:
                fig = go.Figure(go.Bar(
                    x=x2, y=yrg2["Wickets"],
                    marker=dict(color=[GREEN if w >= 3 else oc for w in yrg2["Wickets"]],
                                line=dict(color="white",width=1)),
                    text=yrg2["Wickets"], textposition="outside",
                    textfont=dict(family="Poppins",size=11),
                ))
                pchart(fig, 300, f"Wickets per Year vs {sel_opp}")
                fig.update_yaxes(title_text="Wickets")
                fig.update_xaxes(type="category")
                st.plotly_chart(fig, use_container_width=True)
            with r1b:
                fig2 = go.Figure()
                fig2.add_trace(go.Scatter(
                    x=x2, y=yrg2["Economy"], mode="lines+markers+text",
                    name="Economy", line=dict(color=NAVY,width=2.5),
                    marker=dict(size=10,color=NAVY,line=dict(color="white",width=2)),
                    text=yrg2["Economy"], textposition="top center",
                    textfont=dict(family="Poppins",size=9,color=NAVY)))
                fig2.add_trace(go.Scatter(
                    x=x2, y=yrg2["Bowl_SR"], mode="lines+markers+text",
                    name="Bowl SR", line=dict(color=AMBER,width=2,dash="dot"),
                    marker=dict(size=8,color=AMBER,line=dict(color="white",width=2)),
                    text=yrg2["Bowl_SR"], textposition="bottom center",
                    textfont=dict(family="Poppins",size=9,color=AMBER),
                    yaxis="y2"))
                fig2.update_layout(
                    yaxis2=dict(overlaying="y", side="right",
                        gridcolor="#f0f3f8", linecolor=BORDER,
                        tickfont=dict(size=9)))
                pchart(fig2, 300, f"Economy vs Bowl SR — Year-wise vs {sel_opp}")
                st.plotly_chart(fig2, use_container_width=True)

            r2a, r2b = st.columns(2, gap="large")
            with r2a:
                fig3 = go.Figure()
                fig3.add_trace(go.Bar(name="Good Length", x=x2,
                    y=yrg2["Good"], marker_color=NAVY))
                fig3.add_trace(go.Bar(name="Full Length", x=x2,
                    y=yrg2["Full"], marker_color=AMBER))
                fig3.add_trace(go.Bar(name="Short", x=x2,
                    y=yrg2["Short"], marker_color=STEEL))
                fig3.update_layout(barmode="stack")
                pchart(fig3, 300, f"Wickets by Length — Year-wise vs {sel_opp}")
                fig3.update_yaxes(title_text="Wickets")
                st.plotly_chart(fig3, use_container_width=True,
                                )

            with r2b:
                if "Wicket_RHB" in bw2.columns:
                    hbg = (bw2.groupby("_yr")
                           .agg(RHB=("Wicket_RHB","sum"),
                                LHB=("Wicket_LHB","sum"))
                           .reset_index())
                    fig4 = go.Figure()
                    fig4.add_trace(go.Bar(name="vs RHB", x=hbg["_yr"],
                        y=hbg["RHB"], marker_color=NAVY))
                    fig4.add_trace(go.Bar(name="vs LHB", x=hbg["_yr"],
                        y=hbg["LHB"], marker_color=AMBER))
                    fig4.update_layout(barmode="group")
                    pchart(fig4, 300, f"RHB vs LHB Wickets — Year-wise vs {sel_opp}")
                    fig4.update_yaxes(title_text="Wickets")
                    st.plotly_chart(fig4, use_container_width=True,
                                     )
                else:
                    st.markdown('<div class="no-data">RHB/LHB data unavailable.</div>',
                                unsafe_allow_html=True)

    # ─── ALL-ROUNDER ──────────────────────────────────────────
    else:
        b2_ar  = bat_rows(base2)
        bw2_ar = bowl_rows(base2)
        r1a, r1b = st.columns(2, gap="large")
        with r1a:
            if not b2_ar.empty:
                rg = agg_bat(b2_ar,"_yr").sort_values("_yr")
                fig = go.Figure(go.Bar(x=rg["_yr"],y=rg["Runs"],
                    marker=dict(color=oc,line=dict(color="white",width=1)),
                    text=rg["Runs"],textposition="outside"))
                pchart(fig,300,f"Runs per Year vs {sel_opp}")
                fig.update_yaxes(title_text="Runs")
                fig.update_xaxes(type="category")
                st.plotly_chart(fig,use_container_width=True,  )
        with r1b:
            if not bw2_ar.empty:
                wg = agg_bowl(bw2_ar,"_yr").sort_values("_yr")
                fig2 = go.Figure(go.Bar(x=wg["_yr"],y=wg["Wickets"],
                    marker=dict(color=AMBER,line=dict(color="white",width=1)),
                    text=wg["Wickets"],textposition="outside"))
                pchart(fig2,300,f"Wickets per Year vs {sel_opp}")
                fig2.update_yaxes(title_text="Wickets")
                fig2.update_xaxes(type="category")
                st.plotly_chart(fig2,use_container_width=True,)
        r2a, r2b = st.columns(2, gap="large")
        with r2a:
            if not b2_ar.empty:
                rg2 = agg_bat(b2_ar,"_yr").sort_values("_yr")
                fig3 = go.Figure(go.Scatter(x=rg2["_yr"],y=rg2["Average"],
                    mode="lines+markers",line=dict(color=NAVY,width=2.5),
                    marker=dict(size=10,color=NAVY,line=dict(color="white",width=2))))
                pchart(fig3,300,f"Batting Average per Year vs {sel_opp}")
                fig3.update_yaxes(title_text="Average")
                fig3.update_xaxes(type="category")
                st.plotly_chart(fig3,use_container_width=True)
        with r2b:
            if not bw2_ar.empty:
                eg2 = agg_bowl(bw2_ar,"_yr").sort_values("_yr")
                fig4 = go.Figure(go.Scatter(x=eg2["_yr"],y=eg2["Economy"],
                    mode="lines+markers",line=dict(color=STEEL,width=2.5),
                    marker=dict(size=10,color=STEEL,line=dict(color="white",width=2))))
                pchart(fig4,300,f"Economy per Year vs {sel_opp}")
                fig4.update_yaxes(title_text="Economy")
                fig4.update_xaxes(type="category")
                st.plotly_chart(fig4,use_container_width=True)

# ══════════════════════════════════════════════════════════════
#  SCENARIO 3 — MATCH LEVEL ANALYSIS
# ══════════════════════════════════════════════════════════════
elif scenario == 3:
    oc = opp_clr(sel_opp)

    # ─── BATSMAN ──────────────────────────────────────────────
    if is_bat and not is_all:
        b3 = bat_df.sort_values("Match_date")
        if b3.empty:
            st.markdown('<div class="no-data">No match-level batting data.</div>',
                        unsafe_allow_html=True)
        else:
            xd      = b3["Match_date"].dt.strftime("%d %b")
            runs_v  = list(b3["Runs"].astype(int))
            sr_v    = list(b3["Strike_Rate"].replace(0, float("nan")))
            avg_r   = int(b3["Runs"].mean())

            r1a, r1b = st.columns(2, gap="large")
            with r1a:  # 1. Runs per match — line
                fig = go.Figure()
                fig.add_trace(go.Scatter(
                    x=xd, y=runs_v, mode="lines+markers+text",
                    line=dict(color=oc, width=2.5),
                    marker=dict(
                        size=[14 if r>=100 else (10 if r>=50 else 7) for r in runs_v],
                        color=[AMBER if r>=100 else (GREEN if r>=50 else oc) for r in runs_v],
                        line=dict(color="white",width=2)),
                    text=[str(r) for r in runs_v],
                    textposition="top center",
                    textfont=dict(family="Poppins",size=10),
                    fill="tozeroy", fillcolor="rgba(27,58,107,0.05)"))
                fig.add_hline(y=avg_r, line_dash="dash", line_color=AMBER,
                    annotation_text=f"Avg {avg_r}",
                    annotation_font=dict(color=AMBER,size=9))
                pchart(fig, 300, f"Runs per Match vs {sel_opp}")
                fig.update_xaxes(tickangle=-30, tickfont=dict(size=8))
                fig.update_yaxes(title_text="Runs")
                st.plotly_chart(fig, use_container_width=True,
                                config={"displayModeBar":False})
                

            with r1b:  # 2. Strike Rate per match — bar (color-coded)
                fig2 = go.Figure(go.Bar(
                    x=xd, y=sr_v,
                    marker=dict(
                        color=[GREEN if (s or 0)>=100 else (AMBER if (s or 0)>=70 else RED)
                               for s in sr_v],
                        line=dict(color="white",width=1)),
                    text=[f"{s:.0f}" if s==s else "-" for s in sr_v],
                    textposition="outside",
                    textfont=dict(family="Poppins",size=10),
                ))
                fig2.add_hline(y=100, line_dash="dash", line_color=GREEN,
                    annotation_text="SR 100",
                    annotation_font=dict(color=GREEN,size=9))
                pchart(fig2, 300, f"Strike Rate per Match vs {sel_opp}")
                fig2.update_xaxes(tickangle=-30, tickfont=dict(size=8))
                fig2.update_yaxes(title_text="Strike Rate")
                st.plotly_chart(fig2, use_container_width=True,
                                config={"displayModeBar":False})

            r2a, r2b = st.columns(2, gap="large")
            with r2a:  # 3. Dismissal type — donut
                st.markdown('<div class="chart-card">', unsafe_allow_html=True)
                dis3 = b3[~b3["Dismissed_method"].astype(str).isin(
                    ["Not Applicable","0","nan",""])]
                if not dis3.empty:
                    dm3 = dis3.groupby("Dismissed_method").size().reset_index(name="Count")
                    fig3 = go.Figure(go.Pie(
                        labels=dm3["Dismissed_method"], values=dm3["Count"],
                        hole=0.48,
                        marker=dict(colors=[NAVY,AMBER,STEEL,GREEN,RED],
                                    line=dict(color="white",width=2)),
                        textfont=dict(family="Poppins",size=12),
                        textinfo="label+percent"))
                    fig3.update_layout(
                        paper_bgcolor="rgba(0,0,0,0)", height=300,
                        margin=dict(l=10,r=10,t=44,b=10),
                        title=dict(text=f"Dismissal Types vs {sel_opp}",
                            font=dict(family="Poppins",size=13,color=NAVY),x=0.02),
                        showlegend=False)
                    st.plotly_chart(fig3, use_container_width=True,
                                    config={"displayModeBar":False})
                else:
                    st.markdown('<div class="no-data">No dismissal data.</div>',
                                unsafe_allow_html=True)

            with r2b:  # 4. Dismissal length weakness — bar
                dl3 = b3[~b3["Dismissed_Length"].astype(str).isin(
                    ["Not Applicable","0","nan",""])]
                if not dl3.empty:
                    dlg = (dl3.groupby("Dismissed_Length").size()
                           .reset_index(name="Times")
                           .sort_values("Times", ascending=False))
                    fig4 = go.Figure(go.Bar(
                        x=dlg["Dismissed_Length"], y=dlg["Times"],
                        marker=dict(
                            color=[RED,AMBER,STEEL,NAVY,GREEN][:len(dlg)],
                            line=dict(color="white",width=1)),
                        text=dlg["Times"], textposition="outside",
                        textfont=dict(family="Poppins",size=12),
                    ))
                    pchart(fig4, 300, f"Dismissal Length Weakness vs {sel_opp}")
                    fig4.update_yaxes(title_text="Times Out")
                    st.plotly_chart(fig4, use_container_width=True,
                                    config={"displayModeBar":False})
                else:
                    st.markdown('<div class="no-data">No dismissal length data.</div>',
                                unsafe_allow_html=True)

    # ─── BOWLER ───────────────────────────────────────────────
    elif is_bowl and not is_all:
        bw3 = bowl_df.sort_values("Match_date")
        if bw3.empty:
            st.markdown('<div class="no-data">No match-level bowling data.</div>',
                        unsafe_allow_html=True)
        else:
            xd3   = bw3["Match_date"].dt.strftime("%d %b")
            wkts3 = list(bw3["Wickets"].astype(int))
            eco3  = list(bw3["Economy"].replace(0, float("nan")))
            avg_ec= round(bw3["Economy"].replace(0,float("nan")).mean(), 2)

            r1a, r1b = st.columns(2, gap="large")
            with r1a:  # 1. Wickets per match — bar
                fig = go.Figure(go.Bar(
                    x=xd3, y=wkts3,
                    marker=dict(
                        color=[GREEN if w>=3 else (AMBER if w>=1 else STEEL)
                               for w in wkts3],
                        line=dict(color="white",width=1)),
                    text=wkts3, textposition="outside",
                    textfont=dict(family="Poppins",size=11),
                ))
                pchart(fig, 300, f"Wickets per Match vs {sel_opp}")
                fig.update_xaxes(tickangle=-30, tickfont=dict(size=8))
                fig.update_yaxes(title_text="Wickets")
                st.plotly_chart(fig, use_container_width=True,
                                config={"displayModeBar":False})

            with r1b:  # 2. Economy per match — line
                fig2 = go.Figure()
                fig2.add_trace(go.Scatter(
                    x=xd3, y=eco3, mode="lines+markers",
                    line=dict(color=oc, width=2.5),
                    marker=dict(size=9,
                        color=[RED if (e or 0)>8 else (AMBER if (e or 0)>6 else GREEN)
                               for e in eco3],
                        line=dict(color="white",width=2)),
                    fill="tozeroy", fillcolor="rgba(27,58,107,0.05)"))
                fig2.add_hline(y=avg_ec, line_dash="dash", line_color=AMBER,
                    annotation_text=f"Avg {avg_ec}",
                    annotation_font=dict(color=AMBER,size=9))
                pchart(fig2, 300, f"Economy per Match vs {sel_opp}")
                fig2.update_xaxes(tickangle=-30, tickfont=dict(size=8))
                fig2.update_yaxes(title_text="Economy (RPO)")
                st.plotly_chart(fig2, use_container_width=True,
                                config={"displayModeBar":False})

            r2a, r2b = st.columns(2, gap="large")
            with r2a:  # 3. Wickets by length — stacked bar
                fig3 = go.Figure()
                fig3.add_trace(go.Bar(name="Good Length", x=xd3,
                    y=list(get_col(bw3,"Good_Dissmisals","Good_Dismissals").astype(int)),
                    marker_color=NAVY))
                fig3.add_trace(go.Bar(name="Full Length", x=xd3,
                    y=list(bw3["Full_Dismissals"].astype(int)) if "Full_Dismissals" in bw3 else [0]*len(xd3),
                    marker_color=AMBER))
                fig3.add_trace(go.Bar(name="Short", x=xd3,
                    y=list(get_col(bw3,"Short_Dissmisals","Short_Dismissals").astype(int)),
                    marker_color=STEEL))
                fig3.update_layout(barmode="stack")
                pchart(fig3, 300, f"Wickets by Length per Match vs {sel_opp}")
                fig3.update_xaxes(tickangle=-30, tickfont=dict(size=8))
                fig3.update_yaxes(title_text="Wickets")
                st.plotly_chart(fig3, use_container_width=True,
                                config={"displayModeBar":False})

            with r2b:  # 4. RHB vs LHB match-wise
                if "Wicket_RHB" in bw3.columns:
                    fig4 = go.Figure()
                    fig4.add_trace(go.Bar(name="vs RHB", x=xd3,
                        y=list(bw3["Wicket_RHB"].astype(int)), marker_color=NAVY))
                    fig4.add_trace(go.Bar(name="vs LHB", x=xd3,
                        y=list(bw3["Wicket_LHB"].astype(int)), marker_color=AMBER))
                    fig4.update_layout(barmode="group")
                    pchart(fig4, 300, f"RHB vs LHB per Match vs {sel_opp}")
                    fig4.update_xaxes(tickangle=-30, tickfont=dict(size=8))
                    fig4.update_yaxes(title_text="Wickets")
                    st.plotly_chart(fig4, use_container_width=True,
                                    config={"displayModeBar":False})
                else:
                    st.markdown('<div class="no-data">RHB/LHB data unavailable.</div>',
                                unsafe_allow_html=True)

    # ─── ALL-ROUNDER ──────────────────────────────────────────
    else:
        b3_ar  = bat_df.sort_values("Match_date")
        bw3_ar = bowl_df.sort_values("Match_date")

        r1a, r1b = st.columns(2, gap="large")
        with r1a:  # Bat vs Ball scatter
            mg = pd.merge(
                b3_ar[["Match_date","Runs"]],
                bw3_ar[["Match_date","Wickets"]],
                on="Match_date", how="outer").fillna(0)
            fig = go.Figure(go.Scatter(
                x=mg["Runs"], y=mg["Wickets"], mode="markers",
                marker=dict(size=13, color=oc, opacity=0.82,
                            line=dict(color="white",width=2)),
                text=mg["Match_date"].dt.strftime("%d %b %Y"),
                hovertemplate="<b>%{text}</b><br>Runs: %{x}<br>Wickets: %{y}<extra></extra>"))
            pchart(fig, 300, f"Bat vs Ball Impact — Match Level vs {sel_opp}")
            fig.update_xaxes(title_text="Runs Scored")
            fig.update_yaxes(title_text="Wickets Taken")
            st.plotly_chart(fig, use_container_width=True,
                            config={"displayModeBar":False})

        with r1b:  # Runs line
            if not b3_ar.empty:
                xd_ar = b3_ar["Match_date"].dt.strftime("%d %b")
                rv_ar = list(b3_ar["Runs"].astype(int))
                fig2  = go.Figure()
                fig2.add_trace(go.Scatter(
                    x=xd_ar, y=rv_ar, mode="lines+markers",
                    line=dict(color=NAVY,width=2.5),
                    marker=dict(size=9,color=NAVY,line=dict(color="white",width=2)),
                    fill="tozeroy",fillcolor="rgba(27,58,107,0.05)"))
                pchart(fig2,300,f"Runs per Match vs {sel_opp}")
                fig2.update_xaxes(tickangle=-30,tickfont=dict(size=8))
                fig2.update_yaxes(title_text="Runs")
                st.plotly_chart(fig2,use_container_width=True)

        r2a, r2b = st.columns(2, gap="large")
        with r2a:  # Wickets bar
            if not bw3_ar.empty:
                xb_ar = bw3_ar["Match_date"].dt.strftime("%d %b")
                wv_ar = list(bw3_ar["Wickets"].astype(int))
                fig3  = go.Figure(go.Bar(
                    x=xb_ar, y=wv_ar,
                    marker=dict(color=[GREEN if w>=2 else NAVY for w in wv_ar],
                                line=dict(color="white",width=1)),
                    text=wv_ar,textposition="outside"))
                pchart(fig3,300,f"Wickets per Match vs {sel_opp}")
                fig3.update_xaxes(tickangle=-30,tickfont=dict(size=8))
                fig3.update_yaxes(title_text="Wickets")
                st.plotly_chart(fig3,use_container_width=True)

        with r2b:  # Economy line
            if not bw3_ar.empty:
                xb_ar2 = bw3_ar["Match_date"].dt.strftime("%d %b")
                ec_ar  = list(bw3_ar["Economy"].replace(0,float("nan")))
                fig4   = go.Figure(go.Scatter(
                    x=xb_ar2, y=ec_ar, mode="lines+markers",
                    line=dict(color=STEEL,width=2.5),
                    marker=dict(size=9,
                        color=[RED if (e or 0)>8 else (AMBER if (e or 0)>6 else GREEN) for e in ec_ar],
                        line=dict(color="white",width=2))))
                pchart(fig4,300,f"Economy per Match vs {sel_opp}")
                fig4.update_xaxes(tickangle=-30,tickfont=dict(size=8))
                fig4.update_yaxes(title_text="Economy")
                st.plotly_chart(fig4,use_container_width=True)

# ══════════════════════════════════════════════════════════════
#  ACTION BAR + FOOTER
# ══════════════════════════════════════════════════════════════
yl = f" ({sel_year})" if sel_year != "All" else ""
st.markdown(f"""
<div class="action-bar">
  <div style="display:flex;align-items:center;gap:12px;">
    <a class="act-btn act-report" href="/generate_report" target="_self">
       &nbsp;Generate Report
    </a>
    <span style="color:#8899aa;font-size:.78rem;">
      {sel_player} · {opp_label}{yl} · {total_m} matches
    </span>
  </div>
  <div>
    <a class="act-btn act-compare" href="/compare" target="_self">
       &nbsp;Compare Players
    </a>
  </div>
</div>
""", unsafe_allow_html=True)

st.markdown("""
<div class="footer">
  Cricket Player Performance Analysis  &nbsp;|&nbsp;
  Built with <span>Streamlit</span> 
</div>
""", unsafe_allow_html=True)