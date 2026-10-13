import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import openpyxl
import base64, os

# ─────────────────────────────────────────────
#  PAGE CONFIG
# ─────────────────────────────────────────────
st.set_page_config(
    page_title="Recent Form | Cricket Analysis",
    page_icon="",
    layout="wide",
    initial_sidebar_state="collapsed",
)

# ─────────────────────────────────────────────
#  PATHS
# ─────────────────────────────────────────────
BASE_DIR    = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR    = os.path.dirname(BASE_DIR)
DATASET_DIR = os.path.join(ROOT_DIR, "Dataset")
IMAGE_DIR   = os.path.join(ROOT_DIR, "images")

def _find(fname):
    for d in [DATASET_DIR, ROOT_DIR, BASE_DIR]:
        p = os.path.join(d, fname)
        if os.path.exists(p):
            return p
    return os.path.join(DATASET_DIR, fname)

MASTER_PATH = _find("player_master.csv")
STATS_PATH  = _find("player_stats2.xlsx")

# ─────────────────────────────────────────────
#  GLOBAL CSS  — identical tokens to all pages
# ─────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&family=Poppins:wght@600;700;800&display=swap');

html,body,[class*="css"]{font-family:'Inter',sans-serif;background:#fff !important;}
#MainMenu,footer,header{visibility:hidden;}
section[data-testid="stSidebar"]{display:none !important;}

.block-container,
[data-testid="stAppViewBlockContainer"],
[data-testid="block-container"]{
    padding:0 60px 60px 60px !important;
    max-width:1560px !important;
    margin:0 auto !important;
    background:#fff !important;
}
[data-testid="stAppViewBlockContainer"]>div:first-child{padding-top:0 !important;}

/* ── NAVBAR ── */
.top-header{
    background:#fff;border-bottom:2px solid #1B3A6B;
    padding:0 60px;height:64px;
    display:flex;align-items:center;justify-content:space-between;
    position:fixed;width:100%;top:0;z-index:999;
    box-shadow:0 2px 12px rgba(27,58,107,.08);
    margin:0 -60px 0 -60px;
}
.header-brand{font-family:'Poppins',sans-serif;font-size:1.2rem;font-weight:700;color:#1B3A6B;}
.header-brand span{color:#E8A020;}
.header-nav{display:flex;gap:8px;align-items:center;}
.h-btn{background:#1B3A6B;color:#fff !important;border:none;padding:7px 18px;
       border-radius:6px;font-size:.82rem;font-weight:600;
       text-decoration:none !important;display:inline-block;transition:background .2s;}
.h-btn:hover{background:#243f7a;}
.h-btn-ghost{background:transparent;color:#1B3A6B !important;border:1.5px solid #1B3A6B;
             padding:6px 16px;border-radius:6px;font-size:.82rem;font-weight:600;
             text-decoration:none !important;display:inline-block;transition:background .2s;}
.h-btn-ghost:hover{background:#f0f4fb;}

/* ── PAGE TITLE BAR ── */
.page-title-bar{background:#f7f9fc;border-left:4px solid #1B3A6B;
    padding:14px 60px;margin:40px -60px 28px -60px;}
.page-title-bar h2{font-family:'Poppins',sans-serif;font-size:1.15rem;
    font-weight:700;color:#1B3A6B;margin:0;}
.page-title-bar p{color:#6B7A8D;font-size:.8rem;margin:3px 0 0 0;}

/* ── FILTER PANEL ── */
.filter-wrap{background:#f7f9fc;border:1px solid #dde3ee;border-radius:10px;
             padding:22px 28px;margin-bottom:26px;}
.filter-label{font-size:.75rem;font-weight:600;color:#1B3A6B;
              letter-spacing:.8px;text-transform:uppercase;margin-bottom:6px;}

/* ── WIDGET OVERRIDES ── */
.stSelectbox>div>div{background:#fff !important;border:1.5px solid #c8d0e0 !important;
    border-radius:7px !important;color:#1B3A6B !important;font-size:.88rem !important;}
.stSelectbox>div>div:focus-within{border-color:#1B3A6B !important;
    box-shadow:0 0 0 3px rgba(27,58,107,.1) !important;}
div[data-testid="stButton"]>button{background:#1B3A6B !important;color:#fff !important;
    font-family:'Poppins',sans-serif !important;font-weight:700 !important;
    font-size:.9rem !important;letter-spacing:1px !important;border:none !important;
    border-radius:7px !important;padding:12px 0 !important;width:100% !important;}
div[data-testid="stButton"]>button:hover{background:#243f7a !important;}

/* ── SECTION TITLE ── */
.sec-title{font-family:'Poppins',sans-serif;font-size:1rem;font-weight:700;color:#1B3A6B;
           margin:0 0 14px 0;padding-bottom:8px;border-bottom:2px solid #e8edf5;
           display:flex;align-items:center;gap:8px;}
.sec-title .dot{width:8px;height:8px;background:#1B3A6B;border-radius:50%;display:inline-block;}

/* ── PROFILE CARD ── */
.profile-card{background:#fff;border:1px solid #dde3ee;border-radius:12px;
    padding:28px 22px;box-shadow:0 2px 12px rgba(27,58,107,.07);
    position:relative;overflow:hidden;}
.profile-card::before{content:'';position:absolute;top:0;left:0;right:0;height:4px;
    background:linear-gradient(90deg,#1B3A6B 0%,#E8A020 100%);}
.player-name{font-family:'Poppins',sans-serif;font-size:1.2rem;font-weight:700;
    color:#1B3A6B;text-align:center;margin:12px 0 4px 0;}
.player-role-badge{display:block;width:fit-content;margin:0 auto 18px auto;
    padding:3px 14px;border-radius:20px;font-size:.72rem;font-weight:700;
    letter-spacing:.5px;text-transform:uppercase;}
.badge-batsman{background:#EBF3FF;color:#1B3A6B;border:1px solid #c0d0eb;}
.badge-bowler{background:#FFF0E8;color:#C0501A;border:1px solid #f0c0a0;}
.badge-allrounder{background:#F0F8EE;color:#2A7A2A;border:1px solid #a8d4a0;}
.info-row{display:flex;justify-content:space-between;align-items:center;
          padding:9px 0;border-bottom:1px solid #f0f3f8;font-size:.83rem;}
.info-row:last-child{border-bottom:none;}
.info-lbl{color:#6B7A8D;font-weight:500;}
.info-val{color:#1B3A6B;font-weight:600;}

/* ── MATCH TABLE ── */
.tbl-wrap{background:#fff;border:1px solid #dde3ee;border-radius:12px;
          overflow:hidden;box-shadow:0 2px 10px rgba(27,58,107,.06);}
.match-tbl{width:100%;border-collapse:collapse;}
.match-tbl thead tr{background:#1B3A6B;}
.match-tbl thead th{padding:11px 14px;font-size:.73rem;font-weight:600;
    letter-spacing:.7px;text-transform:uppercase;color:#fff;text-align:left;}
.match-tbl tbody tr{border-bottom:1px solid #f0f3f8;}
.match-tbl tbody tr:last-child{border-bottom:none;}
.match-tbl tbody tr:hover{background:#f7f9fc;}
.match-tbl tbody td{padding:10px 14px;font-size:.83rem;color:#2c3e50;}
.run-val{font-weight:700;color:#1B3A6B;font-size:.95rem;}
.ball-val{color:#8899aa;font-size:.78rem;}
.badge{padding:3px 10px;border-radius:20px;font-size:.7rem;font-weight:700;}
.b-won{background:#EAF6EE;color:#1E7E34;border:1px solid #b8dfc4;}
.b-lose{background:#FDE8E8;color:#C0392B;border:1px solid #f0b8b8;}
.b-nr{background:#F4F4F4;color:#666;border:1px solid #ddd;}

/* ── KPI CARDS ── */
.kpi-card{background:#fff;border:1px solid #dde3ee;border-radius:10px;
    padding:18px 14px 14px 14px;text-align:center;
    box-shadow:0 2px 8px rgba(27,58,107,.05);
    transition:box-shadow .2s,transform .2s;position:relative;overflow:hidden;}
.kpi-card:hover{box-shadow:0 6px 18px rgba(27,58,107,.12);transform:translateY(-2px);}
.kpi-card::after{content:'';position:absolute;bottom:0;left:0;right:0;
    height:3px;background:#1B3A6B;}
.kpi-card.highlight::after{background:#E8A020;}
.kpi-card.hl-g::after{background:#1E7A34;}
.kpi-val{font-family:'Poppins',sans-serif;font-size:1.65rem;font-weight:700;
    color:#1B3A6B;line-height:1;margin-bottom:6px;}
.kpi-lbl{font-size:.7rem;color:#8899aa;font-weight:600;letter-spacing:.6px;text-transform:uppercase;}
.kpi-sub{font-size:.68rem;color:#aab0bc;margin-top:3px;}

/* ── INSIGHT CARDS ── */
.insight-card{background:#f7f9fc;border:1px solid #dde3ee;border-left:4px solid #E8A020;
    border-radius:8px;padding:12px 16px;display:flex;align-items:flex-start;
    gap:12px;margin-bottom:9px;}
.insight-card.pos{border-left-color:#1E7A34;background:#f0f9f2;}
.insight-card.neg{border-left-color:#C0392B;background:#fdf0f0;}
.insight-icon{font-size:1.1rem;flex-shrink:0;margin-top:2px;}
.insight-text{font-size:.84rem;color:#2c3e50;line-height:1.5;}
.insight-text strong{color:#1B3A6B;}

/* ── CHART CARD ── */
.chart-card{background:#fff;border:1px solid #dde3ee;border-radius:12px;
    padding:20px 18px 10px 18px;box-shadow:0 2px 10px rgba(27,58,107,.06);
    margin-bottom:20px;}

/* ── DIVIDER / NO DATA ── */
.divider{border:none;height:1px;background:#e8edf5;margin:28px 0;}
.no-data{text-align:center;padding:50px;color:#8899aa;font-size:.9rem;}
.no-data .icon{font-size:2.5rem;display:block;margin-bottom:10px;}

/* ── ACTION BAR + FOOTER ── */
.action-bar{background:#f7f9fc;border-top:1px solid #dde3ee;
    padding:26px 60px;margin:32px -60px 0 -60px;
    display:flex;align-items:center;justify-content:space-between;}
.act-btn{display:inline-flex;align-items:center;gap:8px;padding:11px 28px;
    border-radius:7px;font-family:'Poppins',sans-serif;font-weight:700;
    font-size:.88rem;letter-spacing:.5px;cursor:pointer;
    text-decoration:none !important;border:none;}
.act-report{background:#fff;color:#1B3A6B !important;border:2px solid #1B3A6B;}
.act-report:hover{background:#EBF3FF;}
.act-compare{background:#1B3A6B;color:#fff !important;}
.act-compare:hover{background:#243f7a;}
.footer{background:#1B3A6B;color:rgba(255,255,255,.7);text-align:center;
    padding:18px 60px;font-size:.8rem;margin:0 -60px -60px -60px;letter-spacing:.3px;}
.footer span{color:#E8A020;font-weight:600;}
</style>
""", unsafe_allow_html=True)


# ─────────────────────────────────────────────
#  DATA LOADING
# ─────────────────────────────────────────────
NON_NUMERIC = {
    "Match_id","Match_date","Opponent","Venue","Country","Match_result",
    "Dismissed_method","Dismissed_Length","Dismissed_against",
    "Dismissed_by","Dismissed _by",
}

@st.cache_data
def load_data():
    master = pd.read_csv(MASTER_PATH)
    master.columns = master.columns.str.strip()
    wb    = openpyxl.load_workbook(STATS_PATH)
    stats = {}
    for sheet in wb.sheetnames:
        ws   = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            stats[sheet.strip()] = pd.DataFrame()
            continue
        cols = [str(c).strip() if c else f"_c{i}" for i, c in enumerate(rows[0])]
        df   = pd.DataFrame(rows[1:], columns=cols)
        df["Match_date"] = pd.to_datetime(df["Match_date"], errors="coerce")
        for c in df.columns:
            if c not in NON_NUMERIC:
                df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0)
        stats[sheet.strip()] = df
    return master, stats

try:
    master_df, stats_dict = load_data()
except Exception as e:
    st.error(f" Data load error: {e}")
    st.stop()


# ─────────────────────────────────────────────
#  HELPERS
# ─────────────────────────────────────────────
NAVY  = "#1B3A6B"
AMBER = "#E8A020"
STEEL = "#4A6FA5"
RED   = "#C0392B"
GREEN = "#1E7A34"

def get_img_b64(filename):
    for d in [IMAGE_DIR, os.path.join(ROOT_DIR, "images"), BASE_DIR]:
        p = os.path.join(d, filename)
        if os.path.exists(p):
            ext  = filename.rsplit(".", 1)[-1].lower()
            mime = "jpeg" if ext in ("jpg","jpeg") else ext
            with open(p, "rb") as f:
                return f"data:image/{mime};base64,{base64.b64encode(f.read()).decode()}"
    return None

def initials_avatar(name):
    ini = "".join(w[0].upper() for w in name.split()[:2])
    return (f'<div style="width:100px;height:100px;border-radius:50%;'
            f'background:#EBF3FF;border:3px solid #1B3A6B;'
            f'display:flex;align-items:center;justify-content:center;'
            f'margin:0 auto;font-family:Poppins,sans-serif;'
            f'font-size:1.8rem;font-weight:700;color:#1B3A6B;">{ini}</div>')

def result_badge(r):
    r = str(r)
    if r == "Won":  return '<span class="badge b-won">Won</span>'
    if r in ("Lose","Loss","Lost"): return '<span class="badge b-lose">Lose</span>'
    return '<span class="badge b-nr">N/R</span>'

def plotly_cfg(fig, height=300, title=""):
    fig.update_layout(
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="#ffffff",
        font=dict(family="Inter", color="#2c3e50", size=11),
        height=height,
        margin=dict(l=8, r=8, t=42, b=8),
        legend=dict(bgcolor="rgba(255,255,255,.9)", bordercolor="#dde3ee",
                    borderwidth=1, font=dict(size=10)),
        title=dict(text=title, font=dict(family="Poppins", size=13, color=NAVY),
                   x=0.02, xanchor="left"),
    )
    fig.update_xaxes(gridcolor="#f0f3f8", linecolor="#dde3ee",
                     tickfont=dict(size=10), zeroline=False)
    fig.update_yaxes(gridcolor="#f0f3f8", linecolor="#dde3ee",
                     tickfont=dict(size=10), zeroline=False)
    return fig

def kpi_html(val, lbl, sub="", cls=""):
    card_cls = f"kpi-card {cls}".strip()
    sub_html = f"<div class='kpi-sub'>{sub}</div>" if sub else ""
    return (f'<div class="{card_cls}">'
            f'<div class="kpi-val">{val}</div>'
            f'<div class="kpi-lbl">{lbl}</div>'
            f'{sub_html}</div>')

def insight(icon, text, variant=""):
    cls = f"insight-card {variant}".strip()
    return (f'<div class="{cls}">'
            f'<span class="insight-icon">{icon}</span>'
            f'<span class="insight-text">{text}</span></div>')

def gcol(df, *names):
    for n in names:
        if n in df.columns:
            return df[n]
    return pd.Series(0, index=df.index)


# ─────────────────────────────────────────────
#  NAVBAR
# ─────────────────────────────────────────────
st.markdown("""
<div class="top-header">
  <div class="header-brand">Cricket Player<span> Performance Analysis</span></div>
  <div class="header-nav">
    <a class="h-btn" href="/userhome" target="_self">Home</a>
    <a class="h-btn" href="/admin" target="_self">Admin</a>
  </div>
</div>
""", unsafe_allow_html=True)

# PAGE TITLE BAR
st.markdown("""
<div class="page-title-bar">
  <h2>Recent Form Analysis — Last 5 Matches</h2>
  <p>Analyze a player's recent performance, form trend and key metrics across their last 5 innings</p>
</div>
""", unsafe_allow_html=True)


# ─────────────────────────────────────────────
#  FILTER PANEL
# ─────────────────────────────────────────────
fc1, fc2,  = st.columns([4, 3])

with fc1:
    st.markdown('<div class="filter-label">Player</div>', unsafe_allow_html=True)
    all_players = sorted(stats_dict.keys())
    sel_player  = st.selectbox("Player", all_players, label_visibility="collapsed")

with fc2:
    st.markdown('<div class="filter-label">Opponent</div>', unsafe_allow_html=True)
    _raw  = stats_dict.get(sel_player, pd.DataFrame())
    _opps = (sorted(_raw["Opponent"].dropna().unique().tolist())
             if not _raw.empty and "Opponent" in _raw.columns else [])
    sel_opp = st.selectbox("Opp", ["All"] + _opps, label_visibility="collapsed")

st.markdown('<div class="filter-label">&nbsp;</div>', unsafe_allow_html=True)
analyze = st.button("ANALYZE NOW", use_container_width=True)



# ─────────────────────────────────────────────
#  SESSION GATE
# ─────────────────────────────────────────────
if analyze:
    st.session_state.update(rf_ok=True, rf_player=sel_player, rf_opp=sel_opp)

if not st.session_state.get("rf_ok"):
    st.markdown("""
    <div class="no-data">
      <span class="icon"></span>
      Select a player &amp; opponent above, then click <strong>Analyze Now</strong>
    </div>
    """, unsafe_allow_html=True)
    st.stop()

sel_player = st.session_state["rf_player"]
sel_opp    = st.session_state["rf_opp"]
df_raw     = stats_dict.get(sel_player, pd.DataFrame())


# ─────────────────────────────────────────────
#  PLAYER META
# ─────────────────────────────────────────────
prow = master_df[master_df["PlayerName"].str.strip() == sel_player]
if not prow.empty:
    pi         = prow.iloc[0]
    p_country  = str(pi.get("Country", "India"))
    p_role     = str(pi.get("Role", "Batsman"))
    p_bat_sty  = str(pi.get("BattingStyle", "RHB"))
    p_bowl_sty = str(pi.get("BowlingStyle", "—"))
    p_dob      = str(pi.get("DOB", "N/A"))
    p_debut    = str(pi.get("DebutYear", "N/A"))
    p_img_file = str(pi.get("ImageFileName", ""))
else:
    p_country = "India"; p_role = "Batsman"; p_bat_sty = "RHB"
    p_bowl_sty = "—";    p_dob = "N/A";    p_debut = "N/A"; p_img_file = ""

rl     = p_role.lower()
is_bat = "batsman"      in rl
is_bowl= "bowler"       in rl
is_ar  = "all-rounder"  in rl or "all rounder" in rl

# Auto-detect from data columns for players not in master CSV
if not (is_bat or is_bowl or is_ar):
    _has_bat  = "Batted_Flag" in df_raw.columns
    _has_bowl = "Bowl_Flag"   in df_raw.columns
    if _has_bat and _has_bowl:
        is_ar   = True;  p_role = "All-Rounder"
    elif _has_bowl:
        is_bowl = True;  p_role = "Bowler"
    else:
        is_bat  = True;  p_role = "Batsman"
    rl = p_role.lower()

badge_cls = ("badge-bowler"     if is_bowl and not is_ar else
             "badge-allrounder" if is_ar                  else
             "badge-batsman")


# ─────────────────────────────────────────────
#  DATA PREP — filter + last 5
# ─────────────────────────────────────────────
if df_raw.empty:
    st.markdown('<div class="no-data"><span class="icon"></span>No data for this player.</div>',
                unsafe_allow_html=True)
    st.stop()

fdf = df_raw.copy()
if sel_opp != "All" and "Opponent" in fdf.columns:
    fdf = fdf[fdf["Opponent"] == sel_opp]

fdf = fdf.sort_values("Match_date", ascending=False)

# Role-specific active rows
if is_bowl and not is_ar:
    active = fdf[fdf["Bowl_Flag"] == 1] if "Bowl_Flag" in fdf.columns else fdf
else:
    active = fdf[fdf["Batted_Flag"] == 1] if "Batted_Flag" in fdf.columns else fdf

last5 = active.head(5).reset_index(drop=True)

if last5.empty:
    st.markdown(
        '<div class="no-data"><span class="icon"></span>'
        'No match data found for the selected filters.</div>',
        unsafe_allow_html=True)
    st.stop()

# Chronological order for charts
last5_chron = last5.iloc[::-1].reset_index(drop=True)


# ═══════════════════════════════════════════════
#  SECTION 1 — PROFILE + TABLE
# ═══════════════════════════════════════════════
col_prof, col_tbl = st.columns([1, 2.4], gap="large")

with col_prof:
    st.markdown('<div class="sec-title"><span class="dot"></span>Player Profile</div>',
                unsafe_allow_html=True)
    img_b64  = get_img_b64(p_img_file) if p_img_file else None
    img_html = (f'<img src="{img_b64}" style="width:100px;height:100px;border-radius:50%;'
                f'object-fit:cover;border:3px solid #1B3A6B;display:block;margin:0 auto;">'
                if img_b64 else initials_avatar(sel_player))

    st.markdown(f"""
    <div class="profile-card">
      {img_html}
      <div class="player-name">{sel_player}</div>
      <div class="player-role-badge {badge_cls}">{p_role}</div>
      <div class="info-row"><span class="info-lbl">Country</span>
        <span class="info-val">🇮🇳 {p_country}</span></div>
      <div class="info-row"><span class="info-lbl">Batting Style</span>
        <span class="info-val">{p_bat_sty}</span></div>
      <div class="info-row"><span class="info-lbl">Bowling Style</span>
        <span class="info-val">{p_bowl_sty}</span></div>
      <div class="info-row"><span class="info-lbl">Date of Birth</span>
        <span class="info-val">{p_dob}</span></div>
      <div class="info-row"><span class="info-lbl">Debut Year</span>
        <span class="info-val">{p_debut}</span></div>
      <div class="info-row"><span class="info-lbl">Opponent Filter</span>
        <span class="info-val">{sel_opp}</span></div>
    </div>
    """, unsafe_allow_html=True)

with col_tbl:
    st.markdown('<div class="sec-title"><span class="dot"></span>'
                'Recent 5 Matches Performance</div>', unsafe_allow_html=True)

    if is_bowl and not is_ar:
        # ── BOWLER TABLE ──────────────────────────
        header = ("<tr><th>Date</th><th>Opponent</th><th>Venue</th>"
                  "<th>Overs</th><th>Figures</th><th>Economy</th><th>Result</th></tr>")
        rows_h = ""
        for _, r in last5.iterrows():
            dt   = r["Match_date"].strftime("%d %b %Y") if pd.notnull(r["Match_date"]) else "N/A"
            bl   = int(r.get("Balls_Bowled", 0))
            ovs  = f"{bl//6}.{bl%6}"
            wk   = int(r.get("Wickets", 0))
            rc   = int(r.get("Runs_Conceded", 0))
            eco  = r.get("Economy", 0)
            eco_s= f"{float(eco):.2f}" if eco else "—"
            res  = str(r.get("Match_result", "N/A"))
            opp  = str(r.get("Opponent", "N/A"))
            ven  = str(r.get("Venue", "N/A"))
            fig  = f"{wk}/{rc}"
            wc   = "color:#1E7A34;font-weight:700;" if wk >= 2 else "color:#2c3e50;font-weight:600;"
            rows_h += (f"<tr><td>{dt}</td><td>{opp}</td><td>{ven}</td>"
                       f"<td style='font-weight:600;color:{NAVY};'>{ovs}</td>"
                       f"<td style='{wc}'>{fig}</td>"
                       f"<td>{eco_s}</td><td>{result_badge(res)}</td></tr>")

    elif is_ar:
        # ── ALL-ROUNDER TABLE ─────────────────────
        header = ("<tr><th>Date</th><th>Opponent</th><th>Venue</th>"
                  "<th>Runs (Balls)</th><th>SR</th>"
                  "<th>Figures</th><th>Economy</th><th>Result</th></tr>")
        rows_h = ""
        bowl_l5 = (fdf[fdf["Bowl_Flag"] == 1].head(5).reset_index(drop=True)
                   if "Bowl_Flag" in fdf.columns else pd.DataFrame())
        for idx, r in last5.iterrows():
            dt    = r["Match_date"].strftime("%d %b %Y") if pd.notnull(r["Match_date"]) else "N/A"
            runs  = int(r.get("Runs", 0))
            balls = int(r.get("Balls_faced", 0))
            sr    = r.get("Strike_Rate", 0)
            sr_s  = f"{float(sr):.1f}" if sr else "—"
            no    = int(r.get("Not_out", 0))
            nm    = "*" if no else ""
            res   = str(r.get("Match_result", "N/A"))
            opp   = str(r.get("Opponent", "N/A"))
            ven   = str(r.get("Venue", "N/A"))
            # Match bowl row by Match_date
            br = bowl_l5[bowl_l5["Match_date"] == r["Match_date"]] if not bowl_l5.empty else pd.DataFrame()
            if not br.empty:
                brow = br.iloc[0]
                wk   = int(brow.get("Wickets", 0))
                rc   = int(brow.get("Runs_Conceded", 0))
                eco  = brow.get("Economy", 0)
                fig  = f"{wk}/{rc}"
                eco_s= f"{float(eco):.2f}" if eco else "—"
            else:
                fig = "DNB"; eco_s = "—"
            rows_h += (f"<tr><td>{dt}</td><td>{opp}</td><td>{ven}</td>"
                       f"<td><span class='run-val'>{runs}{nm}</span>"
                       f" <span class='ball-val'>({balls})</span></td>"
                       f"<td>{sr_s}</td>"
                       f"<td style='font-weight:600;color:{NAVY};'>{fig}</td>"
                       f"<td>{eco_s}</td><td>{result_badge(res)}</td></tr>")

    else:
        # ── BATSMAN TABLE ─────────────────────────
        header = ("<tr><th>Date</th><th>Opponent</th><th>Venue</th>"
                  "<th>Runs (Balls)</th><th>4s</th><th>6s</th>"
                  "<th>SR</th><th>Result</th></tr>")
        rows_h = ""
        for _, r in last5.iterrows():
            dt   = r["Match_date"].strftime("%d %b %Y") if pd.notnull(r["Match_date"]) else "N/A"
            runs = int(r.get("Runs", 0))
            bls  = int(r.get("Balls_faced", 0))
            f4   = int(r.get("4s", 0))
            f6   = int(r.get("6s", 0))
            sr   = r.get("Strike_Rate", 0)
            sr_s = f"{float(sr):.1f}" if sr else "—"
            no   = int(r.get("Not_out", 0))
            nm   = "*" if no else ""
            res  = str(r.get("Match_result", "N/A"))
            opp  = str(r.get("Opponent", "N/A"))
            ven  = str(r.get("Venue", "N/A"))
            rows_h += (f"<tr><td>{dt}</td><td>{opp}</td><td>{ven}</td>"
                       f"<td><span class='run-val'>{runs}{nm}</span>"
                       f" <span class='ball-val'>({bls})</span></td>"
                       f"<td>{f4}</td><td>{f6}</td><td>{sr_s}</td>"
                       f"<td>{result_badge(res)}</td></tr>")

    st.markdown(f"""
    <div class="tbl-wrap">
      <table class="match-tbl">
        <thead>{header}</thead>
        <tbody>{rows_h}</tbody>
      </table>
    </div>""", unsafe_allow_html=True)

st.markdown('<hr class="divider">', unsafe_allow_html=True)


# ═══════════════════════════════════════════════
#  SECTION 2 — KPI CARDS
# ═══════════════════════════════════════════════
st.markdown('<div class="sec-title"><span class="dot"></span>'
            'Key Performance Indicators — Last 5 Matches</div>',
            unsafe_allow_html=True)

wins   = int((last5["Match_result"] == "Won").sum())
losses = int((last5["Match_result"].isin(["Lose","Loss","Lost"])).sum())

# Batting KPIs
if is_bat or is_ar:
    bat_l5      = (last5[last5["Batted_Flag"] == 1]
                   if "Batted_Flag" in last5.columns else last5)
    out_inn     = bat_l5[bat_l5["Not_out"] == 0] if "Not_out" in bat_l5.columns else bat_l5
    total_runs  = int(bat_l5["Runs"].sum())
    avg_runs    = round(bat_l5["Runs"].sum() / max(len(out_inn), 1), 2)
    avg_sr_ser  = bat_l5["Strike_Rate"].replace(0, float("nan"))
    avg_sr      = round(avg_sr_ser.mean(), 1) if not avg_sr_ser.isna().all() else 0
    total_4s    = int(bat_l5["4s"].sum())
    total_6s    = int(bat_l5["6s"].sum())
    fifties     = int(bat_l5["50s"].sum()) if "50s" in bat_l5.columns else 0
    hundreds    = int(bat_l5["100s"].sum()) if "100s" in bat_l5.columns else 0
    dots_bat    = int(bat_l5["Dots"].sum()) if "Dots" in bat_l5.columns else 0
    if not bat_l5.empty:
        hs_row   = bat_l5.loc[bat_l5["Runs"].idxmax()]
        best_bat = f"{int(hs_row['Runs'])}{'*' if int(hs_row.get('Not_out',0)) else ''}"
    else:
        best_bat = "0"

# Bowling KPIs
if is_bowl or is_ar:
    bowl_l5  = (last5[last5["Bowl_Flag"] == 1]
                if "Bowl_Flag" in last5.columns else
                (fdf[fdf["Bowl_Flag"] == 1].head(5).reset_index(drop=True)
                 if "Bowl_Flag" in fdf.columns else pd.DataFrame()))
    if not bowl_l5.empty:
        tot_balls = int(bowl_l5["Balls_Bowled"].sum()) if "Balls_Bowled" in bowl_l5.columns else 0
        overs_str = f"{tot_balls//6}.{tot_balls%6}"
        tot_wkts  = int(bowl_l5["Wickets"].sum())
        eco_ser   = bowl_l5["Economy"].replace(0, float("nan")) if "Economy" in bowl_l5.columns else pd.Series(dtype=float)
        avg_eco   = round(eco_ser.mean(), 2) if not eco_ser.isna().all() else 0.0
        bowl_sr   = round(tot_balls / max(tot_wkts, 1), 1)
        three_fer = int((bowl_l5["Wickets"] >= 3).sum())
        if "Wickets" in bowl_l5.columns and not bowl_l5.empty:
            bf_row   = bowl_l5.sort_values(["Wickets","Runs_Conceded"], ascending=[False,True]).iloc[0]
            best_fig = f"{int(bf_row['Wickets'])}/{int(bf_row.get('Runs_Conceded',0))}"
        else:
            best_fig = "0/0"
    else:
        overs_str = "0.0"; tot_wkts = 0; avg_eco = 0.0
        bowl_sr = 0.0; three_fer = 0; best_fig = "0/0"

# Render KPIs by role
if is_bat and not is_ar:
    kpi_data = [
        (str(total_runs), "Total Runs",    "Last 5 matches",   ""),
        (str(avg_runs),   "Average",       "Per dismissal",    ""),
        (str(avg_sr),     "Strike Rate",   "Avg SR",           ""),
        (f"{total_4s}/{total_6s}", "4s / 6s", "Boundaries",   "highlight"),
        (f"{fifties}/{hundreds}", "50s / 100s", "Milestones",  "highlight"),
        (f"{wins}W-{losses}L",   "Win / Lose", "Results",      ""),
        (best_bat,        "Best Score",    "Top innings",      "highlight"),
    ]
elif is_bowl and not is_ar:
    kpi_data = [
        (overs_str,       "Total Overs",   "Balls bowled",     ""),
        (str(tot_wkts),   "Wickets",       "Last 5 matches",   "highlight"),
        (str(avg_eco),    "Economy",       "Runs per over",    ""),
        (str(bowl_sr),    "Bowl SR",       "Balls per wkt",    ""),
        (str(three_fer),  "3-Wicket Hauls","Big performances", "highlight"),
        (f"{wins}W-{losses}L", "Win / Lose", "Results",        ""),
        (best_fig,        "Best Figures",  "Match haul",       "highlight"),
    ]
else:   # All-Rounder
    kpi_data = [
        (str(total_runs),         "Total Runs",  "Batting",         ""),
        (str(avg_runs),           "Bat Avg",     "Per dismissal",   ""),
        (f"{fifties}/{hundreds}", "50s/100s",    "Milestones",      "highlight"),
        (str(tot_wkts),           "Wickets",     "Bowling",         "highlight"),
        (str(avg_eco),            "Economy",     "Runs per over",   ""),
        (f"{wins}W-{losses}L",    "Win / Lose",  "Results",         ""),
        (best_bat,                "Best Score",  "Highest innings", ""),
    ]

kpi_cols = st.columns(len(kpi_data))
for col_k, (val, lbl, sub, cls) in zip(kpi_cols, kpi_data):
    with col_k:
        st.markdown(kpi_html(val, lbl, sub, cls), unsafe_allow_html=True)

st.markdown('<hr class="divider">', unsafe_allow_html=True)


# ═══════════════════════════════════════════════
#  SECTION 3 — INSIGHTS
# ═══════════════════════════════════════════════
st.markdown('<div class="sec-title"><span class="dot"></span>'
            'Recent Form Insights</div>', unsafe_allow_html=True)

insights_html = ""

if is_bat or is_ar:
    bat_l5 = (last5[last5["Batted_Flag"] == 1]
              if "Batted_Flag" in last5.columns else last5)
    if not bat_l5.empty:
        avg5   = avg_runs
        runs_v = list(bat_l5["Runs"])

        # Avg insight
        insights_html += insight("",
            f"<strong>Average in last 5:</strong> {avg5} runs per dismissal",
            "pos" if avg5 >= 40 else "")

        # 30+ count
        cnt30 = int((bat_l5["Runs"] >= 30).sum())
        insights_html += insight("",
            f"Scored <strong>30+ in {cnt30} of last 5</strong> matches"
            + (" — excellent consistency" if cnt30 >= 4 else
               " — building form" if cnt30 >= 2 else " — form work needed"),
            "pos" if cnt30 >= 4 else ("" if cnt30 >= 2 else "neg"))

        # Strike rate
        insights_html += insight("",
            f"<strong>Strike rate in recent matches:</strong> {avg_sr}"
            + (" — attacking approach " if avg_sr >= 100 else
               " — measured batting" if avg_sr >= 70 else " — below par "),
            "pos" if avg_sr >= 100 else ("" if avg_sr >= 70 else "neg"))

        # Form trend (first 2 vs last 3)
        if len(bat_l5) >= 4:
            sorted_b = bat_l5.sort_values("Match_date")
            early_avg = sorted_b.head(2)["Runs"].mean()
            late_avg  = sorted_b.tail(3)["Runs"].mean()
            if late_avg > early_avg * 1.2:
                trend_lbl = " Improving"; tv = "pos"
            elif late_avg < early_avg * 0.8:
                trend_lbl = " Declining"; tv = "neg"
            else:
                trend_lbl = " Stable"; tv = ""
            insights_html += insight(
                "" if "Improving" in trend_lbl else ("" if "Declining" in trend_lbl else ""),
                f"<strong>Form Trend:</strong> {trend_lbl.split()[-1]} "
                f"(early avg {early_avg:.0f} → recent avg {late_avg:.0f})",
                tv)

        # Boundary reliance
        bnd_runs = total_4s * 4 + total_6s * 6
        bnd_pct  = round(bnd_runs / max(total_runs, 1) * 100, 1)
        insights_html += insight("",
            f"<strong>{bnd_pct}% of runs from boundaries</strong> "
            f"({total_4s} fours, {total_6s} sixes)"
            + (" — high boundary reliance" if bnd_pct > 55 else " — balanced scoring"),
            "" if bnd_pct <= 55 else "neg")

        # Dot-ball pressure
        if dots_bat > 0 and "Balls_faced" in bat_l5.columns:
            tot_balls_f = int(bat_l5["Balls_faced"].sum())
            dot_pct     = round(dots_bat / max(tot_balls_f, 1) * 100, 1)
            insights_html += insight("⏸",
                f"<strong>Dot ball %:</strong> {dot_pct}% of balls faced are dots"
                + (" — under pressure" if dot_pct > 45 else " — rotating strike well"),
                "neg" if dot_pct > 45 else "pos")

if is_bowl or is_ar:
    if not bowl_l5.empty:
        insights_html += insight("",
            f"<strong>Wickets in last 5:</strong> {tot_wkts} wickets at "
            f"an economy of {avg_eco}",
            "pos" if tot_wkts >= 5 else "")

        dp_v = round(
            int(bowl_l5["Dot_Balls"].sum()) / max(int(bowl_l5["Balls_Bowled"].sum()), 1) * 100, 1
        ) if "Dot_Balls" in bowl_l5.columns else 0
        insights_html += insight("",
            f"<strong>Dot ball pressure:</strong> {dp_v}% of deliveries are dots"
            + (" — elite " if dp_v >= 40 else " — building pressure"),
            "pos" if dp_v >= 40 else "")

        if "Economy" in bowl_l5.columns and len(bowl_l5) >= 3:
            eco_trend_early = bowl_l5.sort_values("Match_date").head(2)["Economy"].mean()
            eco_trend_late  = bowl_l5.sort_values("Match_date").tail(3)["Economy"].mean()
            eco_trend       = "Improving " if eco_trend_late < eco_trend_early else "Declining ⚠️"
            tv_eco          = "pos" if "Improving" in eco_trend else "neg"
            insights_html += insight(
                "" if "Improving" in eco_trend else "",
                f"<strong>Bowling trend:</strong> Economy {eco_trend} "
                f"({eco_trend_early:.2f} → {eco_trend_late:.2f})",
                tv_eco)

st.markdown(insights_html, unsafe_allow_html=True)
st.markdown('<hr class="divider">', unsafe_allow_html=True)


# ═══════════════════════════════════════════════
#  SECTION 4 — CHARTS  (4 charts per role)
# ═══════════════════════════════════════════════
st.markdown('<div class="sec-title"><span class="dot"></span>'
            'Performance Charts — Last 5 Matches</div>',
            unsafe_allow_html=True)

# Shared X-axis labels
def make_labels(df):
    return [
        f"M{i+1} vs {str(r['Opponent'])[:3].upper()}"
        for i, r in df.reset_index(drop=True).iterrows()
    ]


# ─── BOWLER CHARTS ────────────────────────────
if is_bowl and not is_ar:
    bowl_ch  = (last5[last5["Bowl_Flag"] == 1]
                if "Bowl_Flag" in last5.columns else last5)
    bowl_chr = bowl_ch.iloc[::-1].reset_index(drop=True)
    xlbls    = make_labels(bowl_chr)

    # Row 1
    cc1, cc2 = st.columns(2, gap="large")

    with cc1:
        # Chart 1 — Wickets Trend (line+markers)
        wk_v = list(bowl_chr["Wickets"].astype(int))
        fig  = go.Figure()
        fig.add_trace(go.Scatter(
            x=xlbls, y=wk_v, mode="lines+markers+text",
            line=dict(color=NAVY, width=2.5),
            marker=dict(size=10, color=[GREEN if w >= 2 else NAVY for w in wk_v],
                        line=dict(color="white", width=2)),
            fill="tozeroy", fillcolor="rgba(27,58,107,0.07)",
            text=wk_v, textposition="top center",
            textfont=dict(family="Poppins", size=12, color=NAVY),
        ))
        plotly_cfg(fig, 300, "Wickets Trend")
        fig.update_yaxes(title_text="Wickets", dtick=1)
        st.plotly_chart(fig, use_container_width=True)

    with cc2:
        # Chart 2 — Economy Trend (column chart)
        eco_v   = [float(v) if v else 0 for v in bowl_chr["Economy"]]
        ec_clrs = [GREEN if e <= 5 else (AMBER if e <= 7 else RED) for e in eco_v]
        fig2    = go.Figure(go.Bar(
            x=xlbls, y=eco_v,
            marker=dict(color=ec_clrs, line=dict(color="white", width=1.5)),
            text=[f"{v:.2f}" for v in eco_v], textposition="outside",
            textfont=dict(family="Poppins", size=11),
        ))
        fig2.add_hline(y=6, line_dash="dash", line_color=AMBER,
                       annotation_text="Eco = 6",
                       annotation_font=dict(color=AMBER, size=10))
        plotly_cfg(fig2, 300, "Economy Rate per Match")
        fig2.update_yaxes(title_text="Economy")
        st.plotly_chart(fig2, use_container_width=True)

    # Row 2
    cc3, cc4 = st.columns(2, gap="large")

    with cc3:
        # Chart 3 — Extras vs Balls Bowled (quadrant-style scatter)
        ext_v   = list(gcol(bowl_chr, "Extras").astype(float))
        balls_v = list(bowl_chr["Balls_Bowled"].astype(float) if "Balls_Bowled" in bowl_chr.columns else [0]*5)
        fig3    = go.Figure()
        fig3.add_trace(go.Scatter(
            x=balls_v, y=ext_v,
            mode="markers+text",
            marker=dict(
                size=14, color=NAVY,
                line=dict(color="white", width=2),
            ),
            text=xlbls, textposition="top center",
            textfont=dict(size=9, color=NAVY),
        ))
        # Quadrant lines
        if balls_v:
            mid_x = sum(balls_v) / len(balls_v)
            mid_y = sum(ext_v)   / max(len(ext_v), 1)
            fig3.add_vline(x=mid_x, line_dash="dot", line_color="#dde3ee")
            fig3.add_hline(y=mid_y, line_dash="dot", line_color="#dde3ee")
        plotly_cfg(fig3, 300, "Extras vs Balls Bowled")
        fig3.update_xaxes(title_text="Balls Bowled")
        fig3.update_yaxes(title_text="Extras")
        st.plotly_chart(fig3, use_container_width=True)

    with cc4:
        # Chart 4 — Dot Ball % Gauge
        tot_b  = int(bowl_ch["Balls_Bowled"].sum()) if "Balls_Bowled" in bowl_ch.columns else 0
        tot_d  = int(gcol(bowl_ch, "Dot_Balls").sum())
        dp_pct = round(tot_d / max(tot_b, 1) * 100, 1)
        fig4   = go.Figure(go.Indicator(
            mode="gauge+number+delta",
            value=dp_pct,
            number=dict(suffix="%", font=dict(family="Poppins", size=36, color=NAVY)),
            delta=dict(reference=40, increasing_color=GREEN, decreasing_color=RED),
            gauge=dict(
                axis=dict(range=[0, 100], tickwidth=1,
                          tickcolor=NAVY, tickfont=dict(size=10)),
                bar=dict(color=NAVY, thickness=0.25),
                bgcolor="white",
                borderwidth=1, bordercolor="#dde3ee",
                steps=[
                    dict(range=[0, 30],  color="#fdf0f0"),
                    dict(range=[30, 50], color="#EBF3FF"),
                    dict(range=[50, 100], color="#f0f9f2"),
                ],
                threshold=dict(line=dict(color=AMBER, width=3),
                               thickness=0.75, value=40),
            ),
            title=dict(text="Dot Ball %", font=dict(family="Poppins", size=13, color=NAVY)),
        ))
        plotly_cfg(fig4, 300, "Dot Balls(last 5 matches)")
        fig4.update_layout(paper_bgcolor="rgba(0,0,0,0)",
                           font=dict(family="Inter"), height=300,
                           margin=dict(l=20, r=20, t=30, b=10))
        st.plotly_chart(fig4, use_container_width=True)
        st.markdown('</div>', unsafe_allow_html=True)


# ─── BATSMAN + ALL-ROUNDER CHARTS ─────────────
else:
    bat_ch  = (last5[last5["Batted_Flag"] == 1]
               if "Batted_Flag" in last5.columns else last5)
    bat_chr = bat_ch.iloc[::-1].reset_index(drop=True)
    xlbls   = make_labels(bat_chr)
    runs_v  = list(bat_chr["Runs"].astype(int))
    sr_v    = [float(v) if float(v) > 0 else None for v in bat_chr["Strike_Rate"]]
    f4_v    = list(bat_chr["4s"].astype(int))
    f6_v    = list(bat_chr["6s"].astype(int))
    dots_v  = list(gcol(bat_chr, "Dots").astype(int))

    avg_v   = sum(r for r in runs_v if r is not None) / max(len(runs_v), 1)

    # Row 1
    cc1, cc2 = st.columns(2, gap="large")

    with cc1:
        # Chart 1 — Runs Trend (line + area)
        fig1 = go.Figure()
        fig1.add_trace(go.Scatter(
            x=xlbls, y=runs_v, mode="lines+markers+text",
            line=dict(color=NAVY, width=2.5),
            marker=dict(size=10, color=NAVY, line=dict(color="white", width=2)),
            fill="tozeroy", fillcolor="rgba(27,58,107,0.07)",
            text=runs_v, textposition="top center",
            textfont=dict(family="Poppins", size=12, color=NAVY),
        ))
        fig1.add_hline(y=avg_v, line_dash="dash", line_color=AMBER,
                       annotation_text=f"Avg {avg_v:.0f}",
                       annotation_font=dict(color=AMBER, size=10))
        plotly_cfg(fig1, 300, "Runs Trend — Last 5 Matches")
        fig1.update_yaxes(title_text="Runs")
        st.plotly_chart(fig1, use_container_width=True)

    with cc2:
        # Chart 2 — Strike Rate Trend (area chart)
        sr_plot  = [v if v else 0 for v in sr_v]
        sr_clrs  = [GREEN if (v or 0) >= 100 else (AMBER if (v or 0) >= 70 else RED) for v in sr_v]
        fig2     = go.Figure()
        fig2.add_trace(go.Scatter(
            x=xlbls, y=sr_plot, mode="lines+markers+text",
            line=dict(color=STEEL, width=2.5),
            marker=dict(size=10, color=sr_clrs, line=dict(color="white", width=2)),
            fill="tozeroy", fillcolor="rgba(74,111,165,0.09)",
            text=[f"{v:.0f}" if v else "0" for v in sr_v],
            textposition="top center",
            textfont=dict(family="Poppins", size=11),
        ))
        fig2.add_hline(y=100, line_dash="dash", line_color=GREEN,
                       annotation_text="SR 100",
                       annotation_font=dict(color=GREEN, size=10))
        plotly_cfg(fig2, 300, "Strike Rate Trend")
        fig2.update_yaxes(title_text="Strike Rate")
        st.plotly_chart(fig2, use_container_width=True)

    # Row 2
    cc3, cc4 = st.columns(2, gap="large")

    with cc3:
        # Chart 3 — Runs vs Team Contribution (lollipop chart)
        # Contribution % = player runs / sum of all bat runs as proxy (no team total col)
        total_runs_all = sum(runs_v) if runs_v else 1
        contrib_pct    = [round(r / max(total_runs_all, 1) * 100, 1) for r in runs_v]
        fig3           = go.Figure()
        # Stems
        for i, (x, y) in enumerate(zip(xlbls, contrib_pct)):
            fig3.add_trace(go.Scatter(
                x=[x, x], y=[0, y],
                mode="lines",
                line=dict(color=NAVY, width=2),
                showlegend=False,
            ))
        # Circles
        fig3.add_trace(go.Scatter(
            x=xlbls, y=contrib_pct,
            mode="markers+text",
            marker=dict(size=16, color=NAVY, line=dict(color=AMBER, width=2.5)),
            text=[f"{v}%" for v in contrib_pct],
            textposition="top center",
            textfont=dict(family="Poppins", size=11, color=NAVY),
            name="Contribution %",
        ))
        plotly_cfg(fig3, 300, "Runs Contribution per Match")
        fig3.update_yaxes(title_text="Contribution %")
        st.plotly_chart(fig3, use_container_width=True)

    with cc4:
        # Chart 4 — Dots vs Boundary Pressure (bubble chart)
        bnd_impact = [f4 * 4 + f6 * 6 for f4, f6 in zip(f4_v, f6_v)]
        sr_bubble  = [float(v) if v else 50 for v in sr_v]
        bubble_sz  = [max(b * 0.8, 8) for b in bnd_impact]

        fig4 = go.Figure()
        fig4.add_trace(go.Scatter(
            x=dots_v, y=bnd_impact,
            mode="markers+text",
            marker=dict(
                size=bubble_sz,
                color=sr_bubble,
                colorscale=[[0, RED], [0.5, AMBER], [1, GREEN]],
                cmin=50, cmax=150,
                colorbar=dict(title="SR", thickness=10,
                              tickfont=dict(size=9)),
                line=dict(color="white", width=1.5),
            ),
            text=xlbls,
            textposition="top center",
            textfont=dict(size=9, color=NAVY),
        ))
        plotly_cfg(fig4, 300, "Dots vs Boundary Impact (Bubble = Boundary Runs, Color = SR)")
        fig4.update_xaxes(title_text="Dot Balls")
        fig4.update_yaxes(title_text="Boundary Runs")
        st.plotly_chart(fig4, use_container_width=True)
        st.markdown('</div>', unsafe_allow_html=True)

    # All-rounder extra bowling row
    if is_ar and "Bowl_Flag" in last5.columns:
        bowl_ch2 = (fdf[fdf["Bowl_Flag"] == 1].head(5).reset_index(drop=True)
                    if "Bowl_Flag" in fdf.columns else pd.DataFrame())
        if not bowl_ch2.empty:
            st.markdown('<div class="sec-title" style="margin-top:10px;">'
                        '<span class="dot"></span>'
                        'Bowling Performance — Last 5 Outings</div>',
                        unsafe_allow_html=True)
            bowl_chr2 = bowl_ch2.iloc[::-1].reset_index(drop=True)
            xlbls2    = make_labels(bowl_chr2)
            cba, cbb  = st.columns(2, gap="large")

            with cba:
                wk2 = list(bowl_chr2["Wickets"].astype(int)) if "Wickets" in bowl_chr2.columns else []
                fb  = go.Figure(go.Bar(
                    x=xlbls2, y=wk2,
                    marker_color=[GREEN if w >= 2 else NAVY for w in wk2],
                    text=wk2, textposition="outside",
                    textfont=dict(family="Poppins", size=12, color=NAVY),
                ))
                plotly_cfg(fb, 280, "Wickets per Outing (Bowling)")
                st.plotly_chart(fb, use_container_width=True)

            with cbb:
                ec2 = [float(v) if v else 0 for v in bowl_chr2["Economy"]] if "Economy" in bowl_chr2.columns else []
                fe  = go.Figure(go.Scatter(
                    x=xlbls2, y=ec2, mode="lines+markers",
                    line=dict(color=NAVY, width=2),
                    marker=dict(size=9, color=NAVY),
                    fill="tozeroy", fillcolor="rgba(27,58,107,0.07)",
                ))
                fe.add_hline(y=6, line_dash="dash", line_color=AMBER)
                plotly_cfg(fe, 280, "Economy Rate Trend (Bowling)")
                st.plotly_chart(fe, use_container_width=True)
                st.markdown('</div>', unsafe_allow_html=True)


# ═══════════════════════════════════════════════
#  ACTION BAR — Generate Report + Compare
# ═══════════════════════════════════════════════
st.markdown(f"""
<div class="action-bar">
  <div>
    <a class="act-btn act-report" href="/generate_report" target="_self">
      &nbsp;Generate Report
    </a>
  </div>
  <div>
    <a class="act-btn act-compare" href="/compare" target="_self">
      &nbsp;Compare Players
    </a>
  </div>
</div>
""", unsafe_allow_html=True)


# ═══════════════════════════════════════════════
#  FOOTER
# ═══════════════════════════════════════════════
st.markdown("""
<div class="footer">
  Cricket Player Performance Analysis Dashboard &nbsp;|&nbsp; Built with Streamlit
</div>
""", unsafe_allow_html=True)